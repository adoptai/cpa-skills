#!/usr/bin/env python3
"""
Build a draft monthly board package from a trial balance and KPI mapping.

What this does, mechanically, and no further:
  * Verifies the trial balance balances (debits = credits) before building
    anything else. This is a hard gate - an unbalanced TB invalidates every
    downstream number, so nothing is built from one.
  * Builds a P&L and Balance Sheet tab by summing TB accounts by
    account_type. Each account's net balance is presented as debit minus
    credit, as recorded - this script does NOT apply revenue/expense or
    asset/liability sign conventions, and does NOT roll equity or retained
    earnings forward. That presentation judgment belongs to the preparer;
    verify sign presentation before issuing the package.
  * Computes each KPI strictly from the account mapping supplied. A KPI
    referencing an account not present in the TB is never estimated - it is
    excluded and listed as "not computed" with the missing account named.
  * Passes flux commentary through verbatim where supplied; marks every
    other line "(commentary not supplied)" rather than inventing narrative.
  * Builds a cash flow tab ONLY if prior-period figures are supplied, as an
    account-level movement rollforward (current - prior per BS account, plus
    net income per the P&L tab). It does NOT classify movements into
    operating / investing / financing activities - that classification is a
    professional judgment call this script does not make. Without prior
    figures, the tab says plainly that cash flow could not be built.

Gates:
  * HARD (blocks the workbook unless --force): the TB does not balance, or
    any TB row has an account_type other than P&L or BS (an unclassifiable
    account cannot be safely placed on either statement).
  * SOFT (reported as exceptions, workbook still produced): a KPI with a
    missing or unmapped account; a flux line item with no commentary
    supplied; a cash-flow account with no prior-period balance.

Usage - emit a blank trial balance worksheet:
    python3 board_package.py --template --out trial_balance_template.csv

Usage - build the package:
    python3 board_package.py --trial-balance tb.csv --kpi-defs kpis.csv \
        --flux flux.csv --prior prior_period.csv \
        --client "Northwind Systems Inc" --period-end 2025-12-31 \
        --out "Northwind - December 2025 Board Package.xlsx"

--trial-balance CSV:
    account, account_type (P&L|BS), debit, credit

--kpi-defs CSV (one row per account that rolls into a KPI):
    kpi_name, account, operator (add|subtract)

--flux CSV (optional):
    line_item, commentary

--prior CSV (optional, prior-period TB balances by account, for the cash
  flow rollforward):
    line_item, prior_amount
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
MONEY = '#,##0.00;[Red](#,##0.00)'

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")

TB_COLS = ["account", "account_type", "debit", "credit"]
KPI_COLS = ["kpi_name", "account", "operator"]
FLUX_COLS = ["line_item", "commentary"]
PRIOR_COLS = ["line_item", "prior_amount"]
ACCOUNT_TYPES = {"P&L", "BS"}
OPERATORS = {"add", "subtract"}


def dec(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def clean(s) -> str:
    return " ".join(str(s or "").split())


def load_csv(path: Path, cols: list[str]) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    out = []
    for i, r in enumerate(rows, start=1):
        x = {k: clean(r.get(k)) for k in cols}
        x["row"] = i
        out.append(x)
    return out


def load_trial_balance(path: Path):
    rows = load_csv(path, TB_COLS)
    if not rows:
        sys.exit(f"No rows in trial balance {path}")
    errors = []
    accounts = []
    for r in rows:
        try:
            debit = dec(r["debit"]) or ZERO
            credit = dec(r["credit"]) or ZERO
        except ValueError as e:
            errors.append(f"TB row {r['row']} ({r['account']}): {e}")
            continue
        atype = r["account_type"]
        if atype not in ACCOUNT_TYPES:
            errors.append(
                f"TB row {r['row']} ({r['account']}): account_type "
                f"{atype!r} is not P&L or BS - cannot classify this "
                f"account onto either statement")
            continue
        accounts.append({
            "account": r["account"], "account_type": atype,
            "debit": debit, "credit": credit, "balance": debit - credit,
        })
    return accounts, errors


def load_kpi_defs(path: Path | None):
    if path is None:
        return defaultdict(list)
    rows = load_csv(path, KPI_COLS)
    by_kpi = defaultdict(list)
    for r in rows:
        op = r["operator"].lower()
        by_kpi[r["kpi_name"]].append({
            "account": r["account"], "operator": op, "row": r["row"],
        })
    return by_kpi


def load_kv(path: Path | None, cols: list[str]):
    if path is None:
        return {}
    rows = load_csv(path, cols)
    return {r[cols[0]]: r[cols[1]] for r in rows}


def write_template(out: Path):
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(TB_COLS)
    print(f"Template written: {out}")


# --------------------------------------------------------------- sheet build

def style_header(ws, row, cols):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(c) + 2)


def sheet_summary(wb, meta, hard_errors, soft_exceptions, pl_total, bs_total, td, tc):
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Board Package Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in meta.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Trial balance totals"])
    ws["A%d" % ws.max_row].font = Font(bold=True)
    ws.append(["Total debits", float(td)])
    ws.append(["Total credits", float(tc)])
    ws.append(["Difference", float(td - tc)])
    for r in range(ws.max_row - 2, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = MONEY
    ws.append([])
    ws.append(["P&L net balance (debit - credit, as recorded)", float(pl_total)])
    ws.append(["Balance Sheet net balance (debit - credit, as recorded)", float(bs_total)])
    for r in range(ws.max_row - 1, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = MONEY
    ws.append([])
    gate_ok = not hard_errors
    cell = ws.cell(row=ws.max_row + 1, column=1,
                    value="GATE: PASSED" if gate_ok else "GATE: FAILED")
    cell.font = OK_FONT if gate_ok else BAD_FONT
    if hard_errors:
        ws.append(["Hard gate failures (blocked the workbook unless --force):"])
        for e in hard_errors:
            ws.append(["", e])
    if soft_exceptions:
        ws.append([])
        ws.append([f"Soft exceptions ({len(soft_exceptions)}) - see Exceptions tab"])


def sheet_statement(wb, title, rows, total):
    ws = wb.create_sheet(title)
    cols = ["account", "debit", "credit", "balance (debit - credit)"]
    style_header(ws, 1, cols)
    r = 2
    for a in rows:
        ws.cell(row=r, column=1, value=a["account"])
        ws.cell(row=r, column=2, value=float(a["debit"])).number_format = MONEY
        ws.cell(row=r, column=3, value=float(a["credit"])).number_format = MONEY
        ws.cell(row=r, column=4, value=float(a["balance"])).number_format = MONEY
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    c = ws.cell(row=r, column=4, value=float(total))
    c.font = Font(bold=True)
    c.number_format = MONEY


def sheet_kpi(wb, kpi_defs, tb_balance, tb_seen):
    ws = wb.create_sheet("KPI")
    cols = ["kpi_name", "value", "components", "status"]
    style_header(ws, 1, cols)
    r = 2
    not_computed = []
    for name, comps in kpi_defs.items():
        missing = [c["account"] for c in comps if c["account"] not in tb_balance]
        bad_op = [c for c in comps if c["operator"] not in OPERATORS]
        if missing or bad_op:
            reasons = []
            if missing:
                reasons.append("account(s) not in TB: " + ", ".join(sorted(set(missing))))
            if bad_op:
                reasons.append("unrecognized operator on row(s): "
                                + ", ".join(str(c["row"]) for c in bad_op))
            not_computed.append((name, "; ".join(reasons)))
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value="(not computed)").fill = WARN_FILL
            ws.cell(row=r, column=3, value=", ".join(c["account"] for c in comps))
            ws.cell(row=r, column=4, value="; ".join(reasons))
            r += 1
            continue
        total = ZERO
        for c in comps:
            v = tb_balance[c["account"]]
            total += v if c["operator"] == "add" else -v
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=float(total)).number_format = MONEY
        ws.cell(row=r, column=3, value=", ".join(
            f"{c['operator']}:{c['account']}" for c in comps))
        ws.cell(row=r, column=4, value="computed")
        r += 1
    return not_computed


def sheet_flux(wb, accounts, flux_map):
    ws = wb.create_sheet("Flux Commentary")
    cols = ["line_item", "commentary"]
    style_header(ws, 1, cols)
    r = 2
    missing = 0
    for a in accounts:
        commentary = flux_map.get(a["account"])
        ws.cell(row=r, column=1, value=a["account"])
        if commentary:
            ws.cell(row=r, column=2, value=commentary)
        else:
            cell = ws.cell(row=r, column=2, value="(commentary not supplied)")
            cell.fill = WARN_FILL
            missing += 1
        r += 1
    return missing


def sheet_cash_flow(wb, bs_accounts, pl_total, prior_map):
    ws = wb.create_sheet("Cash Flow")
    if not prior_map:
        ws.append(["Cash flow could not be built."])
        ws.append(["Reason: no prior-period figures were supplied. Without a prior "
                    "balance per account, account movements cannot be computed, and "
                    "this script does not estimate them."])
        return []
    cols = ["line_item", "prior_amount", "current_amount", "change", "note"]
    style_header(ws, 1, cols)
    r = 2
    no_prior = []
    ws.cell(row=r, column=1, value="Net income (loss) per P&L, as recorded")
    ws.cell(row=r, column=3, value=float(pl_total)).number_format = MONEY
    r += 1
    for a in bs_accounts:
        prior_raw = prior_map.get(a["account"])
        prior_val = dec(prior_raw) if prior_raw not in (None, "") else None
        current = a["balance"]
        ws.cell(row=r, column=1, value=a["account"])
        if prior_val is None:
            no_prior.append(a["account"])
            ws.cell(row=r, column=3, value=float(current)).number_format = MONEY
            ws.cell(row=r, column=4, value=float(current)).number_format = MONEY
            note = ws.cell(row=r, column=5, value="no prior balance - full balance shown as movement")
            note.fill = WARN_FILL
        else:
            change = current - prior_val
            ws.cell(row=r, column=2, value=float(prior_val)).number_format = MONEY
            ws.cell(row=r, column=3, value=float(current)).number_format = MONEY
            ws.cell(row=r, column=4, value=float(change)).number_format = MONEY
        r += 1
    ws.append([])
    ws.append(["Note: this is an account-level movement rollforward only. "
               "Classification into operating / investing / financing activities "
               "is a professional judgment call this script does not make."])
    return no_prior


def sheet_trace(wb, accounts, kpi_defs):
    ws = wb.create_sheet("Trace - Tie-Out")
    cols = ["account", "account_type", "balance", "feeds"]
    style_header(ws, 1, cols)
    feeds_by_account = defaultdict(list)
    for name, comps in kpi_defs.items():
        for c in comps:
            feeds_by_account[c["account"]].append(f"KPI:{name}")
    r = 2
    for a in accounts:
        feeds = [a["account_type"]] + feeds_by_account.get(a["account"], [])
        ws.cell(row=r, column=1, value=a["account"])
        ws.cell(row=r, column=2, value=a["account_type"])
        ws.cell(row=r, column=3, value=float(a["balance"])).number_format = MONEY
        ws.cell(row=r, column=4, value=", ".join(feeds))
        r += 1


def sheet_exceptions(wb, not_computed_kpis, flux_missing, cf_no_prior, dup_accounts):
    ws = wb.create_sheet("Exceptions")
    cols = ["category", "item", "detail"]
    style_header(ws, 1, cols)
    r = 2
    for name, reason in not_computed_kpis:
        ws.cell(row=r, column=1, value="KPI not computed")
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=reason)
        ws.cell(row=r, column=1).fill = WARN_FILL
        r += 1
    for item in flux_missing:
        ws.cell(row=r, column=1, value="Flux commentary not supplied")
        ws.cell(row=r, column=2, value=item)
        r += 1
    for item in cf_no_prior:
        ws.cell(row=r, column=1, value="Cash flow: no prior balance")
        ws.cell(row=r, column=2, value=item)
        r += 1
    for acct, n in dup_accounts.items():
        ws.cell(row=r, column=1, value="Duplicate TB account name")
        ws.cell(row=r, column=2, value=acct)
        ws.cell(row=r, column=3, value=f"appears {n} times - balances summed")
        r += 1
    if r == 2:
        ws.cell(row=2, column=1, value="(no exceptions)")


# ---------------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--trial-balance")
    ap.add_argument("--kpi-defs")
    ap.add_argument("--flux")
    ap.add_argument("--prior")
    ap.add_argument("--client", default="")
    ap.add_argument("--period-end", default="")
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", action="store_true",
                     help="emit a blank trial balance worksheet instead of building")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if args.template:
        write_template(Path(args.out))
        return 0

    if not args.trial_balance or not args.kpi_defs:
        sys.exit("--trial-balance and --kpi-defs are required (or pass --template)")

    accounts, hard_errors = load_trial_balance(Path(args.trial_balance))
    kpi_defs = load_kpi_defs(Path(args.kpi_defs) if args.kpi_defs else None)
    flux_map = load_kv(Path(args.flux) if args.flux else None, FLUX_COLS)
    prior_map = load_kv(Path(args.prior) if args.prior else None, PRIOR_COLS)

    total_debit = sum((a["debit"] for a in accounts), ZERO)
    total_credit = sum((a["credit"] for a in accounts), ZERO)
    if total_debit != total_credit:
        hard_errors.append(
            f"Trial balance does not balance: debits {total_debit:,.2f} vs "
            f"credits {total_credit:,.2f}, difference "
            f"{(total_debit - total_credit):,.2f}")

    tb_balance = {}
    seen_count = defaultdict(int)
    for a in accounts:
        seen_count[a["account"]] += 1
        tb_balance[a["account"]] = tb_balance.get(a["account"], ZERO) + a["balance"]
    dup_accounts = {k: v for k, v in seen_count.items() if v > 1}

    pl_rows = [a for a in accounts if a["account_type"] == "P&L"]
    bs_rows = [a for a in accounts if a["account_type"] == "BS"]
    pl_total = sum((a["balance"] for a in pl_rows), ZERO)
    bs_total = sum((a["balance"] for a in bs_rows), ZERO)

    meta = {
        "Client": args.client or "(not stated)",
        "Period end": args.period_end or "(not stated)",
        "Accounts": len(accounts),
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    valid = not hard_errors

    print("=" * 76)
    print("BOARD PACKAGE BUILDER")
    print("=" * 76)
    print(f"Client {meta['Client']}   Period end {meta['Period end']}")
    print(f"Accounts: {len(accounts)}   Debits {total_debit:,.2f}   "
          f"Credits {total_credit:,.2f}")
    if hard_errors:
        print(f"\nHARD GATE FAILURES ({len(hard_errors)}):")
        for e in hard_errors:
            print(f"  ! {e}")

    if not valid and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("The trial balance does not establish a sound numeric foundation for a "
              "board package. Fix the TB and re-run.")
        print("=" * 76)
        return 1

    wb = Workbook()
    not_computed = []
    flux_missing_list = []
    cf_no_prior = []
    sheet_summary(wb, meta, hard_errors, [], pl_total, bs_total, total_debit, total_credit)
    sheet_statement(wb, "P&L", pl_rows, pl_total)
    sheet_statement(wb, "Balance Sheet", bs_rows, bs_total)
    not_computed = sheet_kpi(wb, kpi_defs, tb_balance, seen_count)
    cf_no_prior = sheet_cash_flow(wb, bs_rows, pl_total, prior_map)
    missing_count = sheet_flux(wb, accounts, flux_map)
    flux_missing_list = [a["account"] for a in accounts if a["account"] not in flux_map]
    sheet_trace(wb, accounts, kpi_defs)
    sheet_exceptions(wb, not_computed, flux_missing_list, cf_no_prior, dup_accounts)
    wb.active = 0

    print(f"\nKPIs not computed: {len(not_computed)}   "
          f"Flux items missing commentary: {missing_count}   "
          f"Cash flow accounts with no prior balance: {len(cf_no_prior)}")

    out = Path(args.out)
    if not valid:
        out = out.with_name(out.stem + " [GATE NOT MET]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'GATE PASSED' if valid else 'GATE NOT MET (forced)'}.  Workbook: {out}")
    return 0 if valid else 1


if __name__ == "__main__":
    sys.exit(main())

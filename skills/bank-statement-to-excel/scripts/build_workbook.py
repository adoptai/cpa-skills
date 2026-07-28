#!/usr/bin/env python3
"""
Prove a normalized statement extract and write the deliverable workbook.

Runs three independent completeness tests and REFUSES to produce a clean
workbook if any of them fail:

  Test A  Balance roll-forward:  opening + credits - debits == closing
  Test B  Statement control totals (if supplied)
  Test C  Line-balance continuity, where the statement prints a running balance

Usage:
    python3 build_workbook.py \
        --transactions work/jan.transactions.csv \
        --opening 42150.22 --closing 38904.71 \
        --out "Acme - First National x4471 - 2026-01 - Transactions.xlsx"

Optional:
    --stmt-deposits 18200.00 --stmt-withdrawals 21445.51   (enables Test B)
    --account-label "First National x4471"
    --period "2026-01-01 to 2026-01-31"
    --force        write the workbook even if a test fails (still marked FAILED)

Input CSV columns (blank where not applicable):
    row_id, source_file, source_page, txn_date, post_date, description,
    check_no, reference, debit, credit, balance, section

Money is handled as Decimal throughout. Nothing is rounded.
"""

from __future__ import annotations

import argparse
import csv
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
MONEY_FMT = '#,##0.00;[Red](#,##0.00)'
DATE_FMT = "yyyy-mm-dd"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
PASS_FONT = Font(bold=True, color="006100")
FAIL_FONT = Font(bold=True, color="9C0006")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")

COLUMNS = [
    "row_id", "source_file", "source_page", "txn_date", "post_date",
    "description", "check_no", "reference", "section",
    "debit", "credit", "balance",
]


def dec(raw) -> Decimal | None:
    """Parse a money string to Decimal. Blank -> None. Never guesses."""
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s.endswith("-"):          # trailing-minus convention
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def parse_date(raw) -> date | None:
    s = str(raw or "").strip()
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        raise ValueError(f"Dates must be ISO YYYY-MM-DD, got {raw!r}")


def load_rows(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {path}")

    out = []
    for i, r in enumerate(rows, start=1):
        d, c = dec(r.get("debit")), dec(r.get("credit"))
        if d is not None and c is not None:
            raise ValueError(
                f"Row {i}: both debit and credit populated. A transaction is one "
                f"or the other. ({r.get('description','')!r})"
            )
        if d is None and c is None:
            raise ValueError(
                f"Row {i}: neither debit nor credit populated. This is usually a "
                f"wrapped description line or a subtotal that should not be a "
                f"transaction row. ({r.get('description','')!r})"
            )
        out.append({
            "row_id": int(r.get("row_id") or i),
            "source_file": (r.get("source_file") or "").strip(),
            "source_page": int(r["source_page"]) if str(r.get("source_page") or "").strip() else None,
            "txn_date": parse_date(r.get("txn_date")),
            "post_date": parse_date(r.get("post_date")),
            "description": " ".join((r.get("description") or "").split()),
            "check_no": (r.get("check_no") or "").strip(),
            "reference": (r.get("reference") or "").strip(),
            "section": (r.get("section") or "").strip(),
            "debit": d, "credit": c,
            "balance": dec(r.get("balance")),
        })
    return out


# --------------------------------------------------------------------------- tests

def test_a(rows, opening: Decimal, closing: Decimal) -> dict:
    deb = sum((r["debit"] for r in rows if r["debit"] is not None), ZERO)
    cre = sum((r["credit"] for r in rows if r["credit"] is not None), ZERO)
    computed = opening + cre - deb
    diff = computed - closing
    return {
        "name": "Test A - Balance roll-forward",
        "detail": "opening + total credits - total debits = statement closing balance",
        "debits": deb, "credits": cre,
        "computed_closing": computed, "stmt_closing": closing,
        "difference": diff, "passed": diff == ZERO,
    }


def test_b(rows, stmt_dep, stmt_wd) -> dict | None:
    if stmt_dep is None and stmt_wd is None:
        return None
    deb = sum((r["debit"] for r in rows if r["debit"] is not None), ZERO)
    cre = sum((r["credit"] for r in rows if r["credit"] is not None), ZERO)
    parts, ok = [], True
    if stmt_dep is not None:
        d = cre - stmt_dep
        ok &= d == ZERO
        parts.append(("Total deposits / credits", cre, stmt_dep, d))
    if stmt_wd is not None:
        d = deb - stmt_wd
        ok &= d == ZERO
        parts.append(("Total withdrawals / debits", deb, stmt_wd, d))
    return {
        "name": "Test B - Statement control totals",
        "detail": "extracted sums agree to the totals printed on the statement",
        "parts": parts, "passed": ok,
    }


def test_c(rows, opening: Decimal) -> dict:
    """Walk the printed running balance. Isolates the exact offending row."""
    with_bal = [r for r in rows if r["balance"] is not None]
    if not with_bal:
        return {
            "name": "Test C - Line-balance continuity",
            "detail": "statement does not print a per-line running balance - not applicable",
            "breaks": [], "passed": True, "applicable": False,
        }
    breaks, prior = [], opening
    for r in rows:
        amt = (r["credit"] or ZERO) - (r["debit"] or ZERO)
        expected = prior + amt
        if r["balance"] is not None:
            if r["balance"] != expected:
                breaks.append({
                    "row_id": r["row_id"], "page": r["source_page"],
                    "date": r["txn_date"], "description": r["description"][:70],
                    "prior": prior, "amount": amt,
                    "expected": expected, "printed": r["balance"],
                    "difference": r["balance"] - expected,
                })
            prior = r["balance"]      # resync so one bad row doesn't cascade
        else:
            prior = expected
    return {
        "name": "Test C - Line-balance continuity",
        "detail": "prior balance +/- amount = printed balance, for every row",
        "breaks": breaks, "passed": not breaks, "applicable": True,
    }


def recomputed_balances(rows, opening: Decimal) -> list[Decimal]:
    out, run = [], opening
    for r in rows:
        run += (r["credit"] or ZERO) - (r["debit"] or ZERO)
        out.append(run)
    return out


# --------------------------------------------------------------------------- excel

def style_header(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def autosize(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def write_transactions(wb, rows, opening) -> None:
    ws = wb.active
    ws.title = "Transactions"
    headers = [
        "Row", "Source File", "Pg", "Txn Date", "Post Date", "Description",
        "Check No", "Reference", "Section", "Debit", "Credit",
        "Balance (printed)", "Balance (recomputed)", "Tie",
    ]
    ws.append(headers)
    style_header(ws, len(headers))

    recomp = recomputed_balances(rows, opening)
    for r, rb in zip(rows, recomp):
        tie = "" if r["balance"] is None else ("OK" if r["balance"] == rb else "BREAK")
        ws.append([
            r["row_id"], r["source_file"], r["source_page"], r["txn_date"],
            r["post_date"], r["description"], r["check_no"], r["reference"],
            r["section"],
            float(r["debit"]) if r["debit"] is not None else None,
            float(r["credit"]) if r["credit"] is not None else None,
            float(r["balance"]) if r["balance"] is not None else None,
            float(rb), tie,
        ])

    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for idx in (4, 5):
            row[idx - 1].number_format = DATE_FMT
        for idx in (10, 11, 12, 13):
            row[idx - 1].number_format = MONEY_FMT
        if row[13].value == "BREAK":
            row[13].font, row[13].fill = FAIL_FONT, FAIL_FILL

    # totals strip
    ws.append([])
    t = ws.max_row + 1
    ws.cell(row=t, column=9, value="TOTALS").font = Font(bold=True)
    ws.cell(row=t, column=10, value=f"=SUBTOTAL(109,J2:J{last})").number_format = MONEY_FMT
    ws.cell(row=t, column=11, value=f"=SUBTOTAL(109,K2:K{last})").number_format = MONEY_FMT
    for c in (10, 11):
        ws.cell(row=t, column=c).font = Font(bold=True)

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:N{last}"
    autosize(ws, {1: 6, 2: 26, 3: 5, 4: 12, 5: 12, 6: 54, 7: 10, 8: 20,
                  9: 20, 10: 14, 11: 14, 12: 17, 13: 19, 14: 8})


def write_proof(wb, results, meta, overall_pass: bool) -> None:
    ws = wb.create_sheet("Proof")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 46

    r = 1
    ws.cell(row=r, column=1, value="COMPLETENESS PROOF").font = Font(bold=True, size=14)
    r += 1
    verdict = ws.cell(
        row=r, column=1,
        value="RESULT: PASS - extraction ties to the statement"
        if overall_pass else
        "RESULT: FAILED - DO NOT RELY ON THIS FILE",
    )
    verdict.font = PASS_FONT if overall_pass else FAIL_FONT
    if not overall_pass:
        verdict.fill = FAIL_FILL
    r += 2

    for k, v in meta.items():
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1

    a = results["a"]
    ws.cell(row=r, column=1, value=a["name"]).font = Font(bold=True, size=12)
    ws.cell(row=r, column=5, value=a["detail"]).font = Font(italic=True, size=9)
    r += 1
    for label, val in [
        ("Opening balance per statement", meta["Opening balance"]),
        ("Add: total credits / deposits", a["credits"]),
        ("Less: total debits / withdrawals", -a["debits"]),
        ("= Computed closing balance", a["computed_closing"]),
        ("Closing balance per statement", a["stmt_closing"]),
        ("Difference", a["difference"]),
    ]:
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=float(val))
        cell.number_format = MONEY_FMT
        if label.startswith(("=", "Difference")):
            cell.font = Font(bold=True)
        if label == "Difference":
            cell.font = PASS_FONT if a["passed"] else FAIL_FONT
            if not a["passed"]:
                cell.fill = FAIL_FILL
        r += 1
    ws.cell(row=r, column=1, value="Test A verdict").font = Font(bold=True)
    vc = ws.cell(row=r, column=2, value="PASS" if a["passed"] else "FAIL")
    vc.font = PASS_FONT if a["passed"] else FAIL_FONT
    r += 2

    b = results.get("b")
    if b:
        ws.cell(row=r, column=1, value=b["name"]).font = Font(bold=True, size=12)
        ws.cell(row=r, column=5, value=b["detail"]).font = Font(italic=True, size=9)
        r += 1
        for h, col in zip(["Item", "Extracted", "Per statement", "Difference"], range(1, 5)):
            hc = ws.cell(row=r, column=col, value=h)
            hc.fill, hc.font = HDR_FILL, HDR_FONT
        r += 1
        for label, extracted, stmt, diff in b["parts"]:
            ws.cell(row=r, column=1, value=label)
            for col, val in ((2, extracted), (3, stmt), (4, diff)):
                c = ws.cell(row=r, column=col, value=float(val))
                c.number_format = MONEY_FMT
            dc = ws.cell(row=r, column=4)
            dc.font = PASS_FONT if diff == ZERO else FAIL_FONT
            if diff != ZERO:
                dc.fill = FAIL_FILL
            r += 1
        ws.cell(row=r, column=1, value="Test B verdict").font = Font(bold=True)
        vc = ws.cell(row=r, column=2, value="PASS" if b["passed"] else "FAIL")
        vc.font = PASS_FONT if b["passed"] else FAIL_FONT
        r += 2
    else:
        ws.cell(row=r, column=1, value="Test B - Statement control totals").font = Font(bold=True, size=12)
        ws.cell(row=r, column=2, value="NOT PERFORMED - statement totals not supplied")
        r += 2

    c = results["c"]
    ws.cell(row=r, column=1, value=c["name"]).font = Font(bold=True, size=12)
    ws.cell(row=r, column=5, value=c["detail"]).font = Font(italic=True, size=9)
    r += 1
    if not c["applicable"]:
        ws.cell(row=r, column=2, value="NOT APPLICABLE - no per-line balance printed")
        r += 1
    elif c["passed"]:
        vc = ws.cell(row=r, column=2, value="PASS - every printed balance agrees")
        vc.font = PASS_FONT
        r += 1
    else:
        ws.cell(row=r, column=1, value=f"{len(c['breaks'])} break(s) - each row below is where "
                                       f"the extract diverges from the statement").font = FAIL_FONT
        r += 1
        heads = ["Row", "Pg", "Date", "Description", "Prior bal", "Amount",
                 "Expected", "Printed", "Difference"]
        for col, h in enumerate(heads, start=1):
            hc = ws.cell(row=r, column=col, value=h)
            hc.fill, hc.font = HDR_FILL, HDR_FONT
        r += 1
        for brk in c["breaks"]:
            ws.cell(row=r, column=1, value=brk["row_id"])
            ws.cell(row=r, column=2, value=brk["page"])
            dc = ws.cell(row=r, column=3, value=brk["date"])
            dc.number_format = DATE_FMT
            ws.cell(row=r, column=4, value=brk["description"])
            for col, key in ((5, "prior"), (6, "amount"), (7, "expected"),
                             (8, "printed"), (9, "difference")):
                mc = ws.cell(row=r, column=col, value=float(brk[key]))
                mc.number_format = MONEY_FMT
            ws.cell(row=r, column=9).font = FAIL_FONT
            r += 1

    r += 1
    ws.cell(row=r, column=1, value="Most likely causes when a test fails, in order:").font = Font(bold=True)
    for line in [
        "1. A subtotal or 'Total Deposits' line captured as a transaction",
        "2. A wrapped description line captured as its own row",
        "3. A transaction on a continuation page missed entirely",
        "4. A debit classified as a credit (or vice versa)",
        "5. An OCR digit misread (0/O, 1/l, 5/S, 8/B) or a dropped leading digit",
        "Test C above identifies the specific row - start there, not with the totals.",
    ]:
        r += 1
        ws.cell(row=r, column=1, value=line)


def write_exceptions(wb, exceptions_csv: Path | None) -> None:
    ws = wb.create_sheet("Exceptions")
    headers = ["#", "Source File", "Pg", "Type", "Description of issue",
               "As presented", "Treatment applied", "Needs preparer decision"]
    ws.append(headers)
    style_header(ws, len(headers))
    n = 0
    if exceptions_csv and exceptions_csv.exists():
        with exceptions_csv.open(newline="", encoding="utf-8-sig") as fh:
            for i, row in enumerate(csv.DictReader(fh), start=1):
                ws.append([i] + [row.get(k, "") for k in
                                 ["source_file", "source_page", "type", "issue",
                                  "as_presented", "treatment", "needs_decision"]])
                n = i
    if n == 0:
        ws.cell(row=2, column=1, value="None")
        ws.cell(row=2, column=5,
                value="No illegible figures, unresolved dates, or ambiguous "
                      "classifications were encountered.")
    autosize(ws, {1: 5, 2: 26, 3: 5, 4: 20, 5: 56, 6: 18, 7: 34, 8: 22})


def write_source_map(wb, rows) -> None:
    ws = wb.create_sheet("Source Map")
    headers = ["Source File", "Pg", "Transactions", "Debits on page",
               "Credits on page", "Net"]
    ws.append(headers)
    style_header(ws, len(headers))

    agg: dict[tuple, dict] = {}
    for r in rows:
        key = (r["source_file"], r["source_page"])
        a = agg.setdefault(key, {"n": 0, "d": ZERO, "c": ZERO})
        a["n"] += 1
        a["d"] += r["debit"] or ZERO
        a["c"] += r["credit"] or ZERO

    for (f, p), a in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1] or 0)):
        ws.append([f, p, a["n"], float(a["d"]), float(a["c"]), float(a["c"] - a["d"])])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for idx in (4, 5, 6):
            row[idx - 1].number_format = MONEY_FMT
    ws.freeze_panes = "A2"
    autosize(ws, {1: 30, 2: 6, 3: 14, 4: 16, 5: 16, 6: 16})


# --------------------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--transactions", required=True)
    ap.add_argument("--opening", required=True)
    ap.add_argument("--closing", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--stmt-deposits")
    ap.add_argument("--stmt-withdrawals")
    ap.add_argument("--exceptions", help="optional exceptions CSV")
    ap.add_argument("--account-label", default="")
    ap.add_argument("--period", default="")
    ap.add_argument("--force", action="store_true",
                    help="write the workbook even if a test fails")
    args = ap.parse_args()

    rows = load_rows(Path(args.transactions))
    opening, closing = dec(args.opening), dec(args.closing)
    if opening is None or closing is None:
        sys.exit("--opening and --closing are required and must be numeric.")

    results = {
        "a": test_a(rows, opening, closing),
        "b": test_b(rows, dec(args.stmt_deposits), dec(args.stmt_withdrawals)),
        "c": test_c(rows, opening),
    }
    overall = results["a"]["passed"] and results["c"]["passed"] and (
        results["b"]["passed"] if results["b"] else True
    )

    meta = {
        "Account": args.account_label or "(not supplied)",
        "Statement period": args.period or "(not supplied)",
        "Source file(s)": ", ".join(sorted({r["source_file"] for r in rows if r["source_file"]})) or "(not supplied)",
        "Transactions extracted": len(rows),
        "Opening balance": opening,
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    # ---- console report
    print("=" * 68)
    print("COMPLETENESS PROOF")
    print("=" * 68)
    a = results["a"]
    print(f"Transactions extracted : {len(rows)}")
    print(f"Opening balance        : {opening:>14,.2f}")
    print(f"Total credits          : {a['credits']:>14,.2f}")
    print(f"Total debits           : {a['debits']:>14,.2f}")
    print(f"Computed closing       : {a['computed_closing']:>14,.2f}")
    print(f"Statement closing      : {a['stmt_closing']:>14,.2f}")
    print(f"Difference             : {a['difference']:>14,.2f}   "
          f"{'PASS' if a['passed'] else '*** FAIL ***'}")
    if results["b"]:
        print(f"\nTest B (control totals): "
              f"{'PASS' if results['b']['passed'] else '*** FAIL ***'}")
        for label, ex, st, df in results["b"]["parts"]:
            print(f"  {label:<28} extracted {ex:>12,.2f}  stmt {st:>12,.2f}  diff {df:>10,.2f}")
    c = results["c"]
    if not c["applicable"]:
        print("\nTest C (line continuity): N/A - no per-line balance printed")
    elif c["passed"]:
        print("\nTest C (line continuity): PASS")
    else:
        print(f"\nTest C (line continuity): *** FAIL - {len(c['breaks'])} break(s) ***")
        for brk in c["breaks"][:15]:
            print(f"  row {brk['row_id']:>4} pg {brk['page']}  {brk['date']}  "
                  f"{brk['description'][:40]:<40} expected {brk['expected']:>12,.2f} "
                  f"printed {brk['printed']:>12,.2f}  diff {brk['difference']:>10,.2f}")
        if len(c["breaks"]) > 15:
            print(f"  ... and {len(c['breaks']) - 15} more")

    if not overall and not args.force:
        print("\n" + "=" * 68)
        print("WORKBOOK NOT WRITTEN. The extract does not tie to the statement.")
        print("Fix the extract and re-run. Do not deliver an unproven file.")
        print("Check, in order: subtotals captured as transactions; wrapped")
        print("description lines; missed continuation pages; debit/credit flipped;")
        print("OCR digit misreads. Test C above names the offending row.")
        print("Use --force only to produce a marked-FAILED file for debugging.")
        print("=" * 68)
        return 1

    wb = Workbook()
    write_transactions(wb, rows, opening)
    write_proof(wb, results, meta, overall)
    write_exceptions(wb, Path(args.exceptions) if args.exceptions else None)
    write_source_map(wb, rows)
    wb.active = 0

    out = Path(args.out)
    if not overall:
        out = out.with_name(out.stem + " [FAILED PROOF]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    print("\n" + ("PROOF PASSED. " if overall else "PROOF FAILED (forced). "))
    print(f"Workbook: {out}")
    return 0 if overall else 1


if __name__ == "__main__":
    sys.exit(main())

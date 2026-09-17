#!/usr/bin/env python3
"""
Balance sheet reconciliation pack, built account by account across the full
chart rather than one account at a time.

For every balance sheet account this proves the same thing bank-rec-to-gl
proves for cash alone: GL balance = supporting schedule balance + itemized,
dated reconciling items, with nothing left unexplained. An account is not
"reconciled" here because someone typed the word - it is reconciled because
the math ties AND a preparer and a reviewer are both on record.

Gates:
  * every balance sheet account in the trial balance appears in the pack -
    reconciled, exception, tied-but-unsigned, or not received. None dropped.
  * unexplained = GL balance - (schedule balance + reconciling items). Any
    account where that is not exactly zero is an EXCEPTION and stays visible.
    Never plugged, never rounded away.
  * nothing is marked RECONCILED without both a preparer and a reviewer name
    and date on record. A tie with no signature is progress, not a reconciled
    account.
  * the pack is COMPLETE only when every balance sheet account is RECONCILED.
    Otherwise it is produced (so preparers can keep working from it) but
    reported and filed as incomplete, and the script exits non-zero.

What this script does NOT do: it does not decide whether an account should be
reconciled this period at all (that is a close-checklist waiver decision -
see month-end-close-checklist), it does not classify what a reconciling item
IS (deposit in transit vs. book error vs. timing - that is judgment, and for
cash specifically it is bank-rec-to-gl's job), and it does not choose
materiality.

Usage:
    python3 recon_pack.py \
        --trial-balance tb.csv --schedules schedules.csv --items items.csv \
        --prior-pack prior_items.csv --materiality 5000 --as-of 2025-12-31 \
        --client "Halstead Regional Foods" --period "2025-12" \
        --out "Halstead - 2025-12 Balance Sheet Recon Pack.xlsx"

--trial-balance CSV (same schema trial-balance-integrity / month-end-close-checklist use):
    account, name, debit, credit, type, contra (yes/no)
      type: asset | contra_asset | liability | contra_liability | equity
            (revenue/income/expense/cogs rows are skipped - not balance sheet)

--schedules CSV, one row per account with a supporting schedule prepared:
    account, schedule_balance, source, prepared_by, prepared_date,
    reviewed_by, reviewed_date, notes
      schedule_balance is in the same sign convention as the trial balance's
      net (debit - credit) for that account.

--items CSV, the itemized reconciling items explaining GL vs. schedule
(zero or more rows per account):
    account, item_date, description, amount, notes
      amount is signed: the sign that, added to the schedule balance, moves
      it toward the GL balance.

--prior-pack CSV, this account's UNRESOLVED items as of last period's pack
(for recurrence detection only):
    account, description, amount, item_date, consecutive_periods

Money is Decimal throughout. Nothing is rounded.
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
MONEY = '#,##0.00;[Red](#,##0.00)'
DATEF = "yyyy-mm-dd"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
TOP = Border(top=Side(style="thin"))

BS_TYPES = {"asset", "contra_asset", "liability", "contra_liability", "equity"}
IS_TYPES = {"revenue", "income", "expense", "cogs", "cost_of_goods_sold"}

STATUS_ORDER = ("RECONCILED", "PREPARED - NOT REVIEWED", "TIED - NOT SIGNED",
                "EXCEPTION", "NOT RECEIVED")

# Which skill in this repository typically produces the supporting schedule
# for an account like this. Advisory only - never used to compute anything.
SCHEDULE_SOURCE_HINTS = [
    (("cash", "bank", "checking", "money market", "operating account", "petty"),
     "bank-rec-to-gl (or bank-rec-review if someone else prepared it)"),
    (("receivable", "a/r", "unbilled", "contract asset", "allowance", "doubtful"),
     "ar-aging-tie-out"),
    (("fixed asset", "machinery", "equipment", "property", "plant", "leasehold",
      "vehicle", "furniture", "accumulated depreciation", "cip",
      "construction in progress"),
     "depreciation-tie-out"),
    (("payable", "a/p", "accrued", "accrual"),
     "cutoff-and-unrecorded-liabilities"),
    (("wage", "salar", "payroll", "fica", "medicare", "withholding", "941"),
     "payroll-tax-reconciliation"),
    (("sales tax", "use tax", "vat", "gst"),
     "sales-tax-reconciliation"),
    (("suspense", "clearing", "conversion", "unallocated", "temporary", "holding"),
     "trial-balance-integrity (this should be empty - investigate, do not schedule around it)"),
    (("equity", "capital", "retained earnings", "distribution", "dividend",
      "common stock", "member"),
     "trial-balance-integrity (rollforward vs. net income and distributions)"),
]


# ------------------------------------------------------------------ parsing

def dec(raw) -> Decimal | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def d0(raw) -> Decimal:
    v = dec(raw)
    return v if v is not None else ZERO


def clean(s) -> str:
    return " ".join(str(s or "").split())


def pdate(raw) -> date | None:
    s = clean(raw)
    if not s:
        return None
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s.split(" ")[0], f).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date {raw!r}. Prefer ISO YYYY-MM-DD.")


def rows_of(path: Path, label: str) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {label} ({path})")
    return rows


def schedule_hint(name: str) -> str:
    low = name.lower()
    for words, hint in SCHEDULE_SOURCE_HINTS:
        if any(w in low for w in words):
            return hint
    return "(no repository skill matched - prepare the schedule manually)"


# ------------------------------------------------------------------ loading

def load_tb(path: Path) -> list[dict]:
    out = []
    for i, r in enumerate(rows_of(path, "trial balance"), start=1):
        acct = clean(r.get("account"))
        if not acct:
            sys.exit(f"Trial balance row {i}: no account number. Fix the source export.")
        t = clean(r.get("type")).lower()
        if t in IS_TYPES:
            continue
        if t not in BS_TYPES:
            sys.exit(
                f"Trial balance row {i} ({acct}): type '{t or '(blank)'}' is not one of "
                f"{sorted(BS_TYPES)} or an income-statement type {sorted(IS_TYPES)}. "
                f"Run trial-balance-integrity first to get every account typed."
            )
        dr, cr = d0(r.get("debit")), d0(r.get("credit"))
        name = clean(r.get("name"))
        out.append({
            "account": acct,
            "name": name,
            "type": t,
            "contra": clean(r.get("contra")).lower() in ("yes", "y", "true", "1"),
            "gl_net": dr - cr,
            "hint": schedule_hint(name),
            "schedule_balance": None,
            "source": "", "prepared_by": "", "prepared_date": None,
            "reviewed_by": "", "reviewed_date": None, "sched_notes": "",
            "items": [],
            "items_total": ZERO,
            "raw_difference": None,
            "unexplained": None,
            "status": "",
            "escalate": False,
        })
    if not out:
        sys.exit("No balance sheet accounts found (asset/liability/equity). "
                 "Check the trial balance's 'type' column.")
    seen = set()
    for a in out:
        if a["account"] in seen:
            sys.exit(f"Duplicate account number '{a['account']}' in the trial balance. "
                     f"Run trial-balance-integrity to resolve duplicates before this pack.")
        seen.add(a["account"])
    return out


def load_schedules(path: Path | None) -> dict[str, dict]:
    out = {}
    if not path or not path.exists():
        return out
    for r in rows_of(path, "schedules"):
        acct = clean(r.get("account"))
        if not acct:
            continue
        bal = dec(r.get("schedule_balance"))
        if bal is None:
            sys.exit(f"Schedule row for account {acct}: schedule_balance is required "
                     f"and must be numeric. A blank schedule balance is not a zero "
                     f"balance - state it explicitly or omit the row.")
        out[acct] = {
            "schedule_balance": bal,
            "source": clean(r.get("source")),
            "prepared_by": clean(r.get("prepared_by")),
            "prepared_date": pdate(r.get("prepared_date")),
            "reviewed_by": clean(r.get("reviewed_by")),
            "reviewed_date": pdate(r.get("reviewed_date")),
            "notes": clean(r.get("notes")),
        }
    return out


def load_items(path: Path | None) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = defaultdict(list)
    if not path or not path.exists():
        return out
    for i, r in enumerate(rows_of(path, "items"), start=1):
        acct = clean(r.get("account"))
        if not acct:
            sys.exit(f"Item row {i}: no account number.")
        amt = dec(r.get("amount"))
        if amt is None:
            sys.exit(f"Item row {i} (account {acct}): amount is required and must be "
                     f"numeric and signed. An item with no amount cannot explain "
                     f"anything - it is not a reconciling item.")
        d = pdate(r.get("item_date"))
        if d is None:
            sys.exit(f"Item row {i} (account {acct}): item_date is required - an "
                     f"undated item cannot be aged, and age is the whole point of "
                     f"carrying an open item forward.")
        out[acct].append({
            "account": acct, "item_date": d,
            "description": clean(r.get("description")) or "(no description)",
            "amount": amt, "notes": clean(r.get("notes")),
        })
    return out


def load_prior(path: Path | None) -> dict[tuple, int]:
    out = {}
    if not path or not path.exists():
        return out
    for r in rows_of(path, "prior pack"):
        acct = clean(r.get("account"))
        desc = clean(r.get("description"))
        amt = dec(r.get("amount"))
        if not acct or amt is None:
            continue
        n = int(d0(r.get("consecutive_periods"))) or 1
        out[(acct, desc, amt)] = n
    return out


# ------------------------------------------------------------------ logic

def bucket(days: int) -> str:
    for lim, name in ((30, "0-30"), (60, "31-60"), (90, "61-90"), (180, "91-180")):
        if days <= lim:
            return name
    return "180+"


def build_pack(tb: list[dict], schedules: dict, items: dict, prior: dict,
              materiality: Decimal | None, asof: date) -> list[str]:
    """Annotate each account in place. Returns a list of orphan warnings."""
    warnings: list[str] = []
    tb_accounts = {a["account"] for a in tb}

    for acct in schedules:
        if acct not in tb_accounts:
            warnings.append(f"schedules.csv has account '{acct}' which is not a "
                            f"balance sheet account on the trial balance (typo, or "
                            f"an income-statement account was scheduled by mistake).")
    for acct in items:
        if acct not in tb_accounts:
            warnings.append(f"items.csv has account '{acct}' which is not a balance "
                            f"sheet account on the trial balance.")

    for a in tb:
        sch = schedules.get(a["account"])
        acct_items = items.get(a["account"], [])
        for it in acct_items:
            age = (asof - it["item_date"]).days
            it["age_days"] = age
            it["bucket"] = bucket(age)
            key = (a["account"], it["description"], it["amount"])
            prior_n = prior.get(key)
            it["recurring"] = prior_n is not None
            it["consecutive_periods"] = (prior_n + 1) if prior_n else 1
        a["items"] = sorted(acct_items, key=lambda x: x["item_date"])
        a["items_total"] = sum((it["amount"] for it in acct_items), ZERO)

        if sch is None:
            a["status"] = "NOT RECEIVED"
            a["escalate"] = materiality is None or abs(a["gl_net"]) >= materiality
            continue

        a["schedule_balance"] = sch["schedule_balance"]
        a["source"] = sch["source"]
        a["prepared_by"] = sch["prepared_by"]
        a["prepared_date"] = sch["prepared_date"]
        a["reviewed_by"] = sch["reviewed_by"]
        a["reviewed_date"] = sch["reviewed_date"]
        a["sched_notes"] = sch["notes"]

        raw_diff = a["gl_net"] - a["schedule_balance"]
        unexplained = raw_diff - a["items_total"]
        a["raw_difference"] = raw_diff
        a["unexplained"] = unexplained

        if unexplained != ZERO:
            a["status"] = "EXCEPTION"
            a["escalate"] = materiality is None or abs(unexplained) >= materiality
        elif not a["prepared_by"] or not a["prepared_date"]:
            a["status"] = "TIED - NOT SIGNED"
        elif not a["reviewed_by"] or not a["reviewed_date"]:
            a["status"] = "PREPARED - NOT REVIEWED"
        else:
            a["status"] = "RECONCILED"

    return warnings


# ------------------------------------------------------------------ workbook

def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


def status_font(status: str) -> Font:
    if status == "RECONCILED":
        return OK_FONT
    if status in ("EXCEPTION", "NOT RECEIVED"):
        return BAD_FONT
    return Font(bold=True, color="9C6500")  # in-progress states


def sheet_summary(wb, tb, meta, materiality, complete: bool):
    ws = wb.active
    ws.title = "Pack Summary"
    widths(ws, {1: 4, 2: 56, 3: 18, 4: 18, 5: 40})
    r = 1

    def line(label, val=None, *, bold=False, note=""):
        nonlocal r
        c = ws.cell(row=r, column=2, value=label)
        c.font = Font(bold=bold)
        if val is not None:
            m = ws.cell(row=r, column=3, value=val)
            m.font = Font(bold=bold)
        if note:
            ws.cell(row=r, column=5, value=note).font = Font(italic=True, size=9)
        r += 1

    ws.cell(row=r, column=2, value="BALANCE SHEET RECONCILIATION PACK").font = \
        Font(bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    v = ws.cell(row=r, column=2,
               value="COMPLETE - every balance sheet account is reconciled and signed"
               if complete else
               "INCOMPLETE - see accounts below that are not yet reconciled")
    v.font = OK_FONT if complete else BAD_FONT
    if not complete:
        v.fill = BAD_FILL
    r += 2

    counts = defaultdict(int)
    dollars = defaultdict(lambda: ZERO)
    for a in tb:
        counts[a["status"]] += 1
        dollars[a["status"]] += abs(a["gl_net"])

    c = ws.cell(row=r, column=2, value="COVERAGE PROOF")
    c.font, c.fill = Font(bold=True, size=12), SUB_FILL
    r += 1
    for s in STATUS_ORDER:
        line(f"  {s}", None)
        ws.cell(row=r - 1, column=3, value=counts.get(s, 0))
        ws.cell(row=r - 1, column=4, value=float(dollars.get(s, ZERO))).number_format = MONEY
    total_line_row = r
    ws.cell(row=r, column=2, value="  Total balance sheet accounts").font = Font(bold=True)
    ws.cell(row=r, column=3, value=len(tb)).font = Font(bold=True)
    ws.cell(row=r, column=3).border = TOP
    r += 2

    c = ws.cell(row=r, column=2, value="THE GATE")
    c.font, c.fill = Font(bold=True, size=12), SUB_FILL
    r += 1
    for g in [
        "1. Every balance sheet account in the trial balance appears above - "
        "reconciled, exception, tied-but-unsigned, or not received. None dropped.",
        "2. Unexplained = GL balance - (schedule balance + reconciling items). Any "
        "account where that is not exactly zero is an EXCEPTION, shown, never plugged.",
        "3. RECONCILED requires both a preparer and a reviewer name and date on "
        "record - a tie with no signature is TIED - NOT SIGNED, not reconciled.",
        "4. The pack is COMPLETE only when every account above is RECONCILED.",
    ]:
        ws.cell(row=r, column=2, value=g).alignment = Alignment(wrap_text=True)
        r += 1
    r += 1

    if materiality is not None:
        line("Materiality (for the ESCALATE flag only - not for the gate)",
             float(materiality))
        ws.cell(row=r - 1, column=3).number_format = MONEY
        r += 1


def sheet_chart(wb, tb):
    ws = wb.create_sheet("Chart Status")
    heads = ["Account", "Name", "Type", "GL Balance", "Schedule Balance",
             "Reconciling Items", "Unexplained", "Status", "Escalate",
             "Prepared By", "Prepared Date", "Reviewed By", "Reviewed Date",
             "Source", "Suggested Skill", "Notes"]
    ws.append(heads)
    hdr(ws, len(heads))
    for a in sorted(tb, key=lambda x: (STATUS_ORDER.index(x["status"]), x["account"])):
        ws.append([
            a["account"], a["name"], a["type"], float(a["gl_net"]),
            float(a["schedule_balance"]) if a["schedule_balance"] is not None else None,
            float(a["items_total"]) if a["schedule_balance"] is not None else None,
            float(a["unexplained"]) if a["unexplained"] is not None else None,
            a["status"], "YES" if a["escalate"] else "",
            a["prepared_by"], a["prepared_date"], a["reviewed_by"], a["reviewed_date"],
            a["source"], a["hint"], a["sched_notes"],
        ])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col in (4, 5, 6, 7):
            if row[col - 1].value is not None:
                row[col - 1].number_format = MONEY
        row[7].font = status_font(row[7].value)
        if row[7].value in ("EXCEPTION", "NOT RECEIVED"):
            row[7].fill = BAD_FILL
        elif row[7].value != "RECONCILED":
            row[7].fill = WARN_FILL
        if row[8].value == "YES":
            row[8].font, row[8].fill = BAD_FONT, BAD_FILL
        for col in (10, 12):
            if row[col - 1].value:
                row[col - 1].number_format = DATEF
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:P{max(ws.max_row, 2)}"
    widths(ws, {1: 10, 2: 30, 3: 14, 4: 14, 5: 14, 6: 14, 7: 13, 8: 20, 9: 9,
               10: 14, 11: 13, 12: 14, 13: 13, 14: 22, 15: 42, 16: 34})


def sheet_items(wb, tb, asof: date):
    ws = wb.create_sheet("Reconciling Items")
    heads = ["Account", "Name", "Item Date", "Age (d)", "Bucket", "Description",
             "Amount", "Recurring", "Consecutive Periods", "Notes"]
    ws.append(heads)
    hdr(ws, len(heads))
    n = 0
    for a in sorted(tb, key=lambda x: x["account"]):
        for it in a["items"]:
            n += 1
            ws.append([
                a["account"], a["name"], it["item_date"], it["age_days"], it["bucket"],
                it["description"], float(it["amount"]),
                "YES" if it["recurring"] else "", it["consecutive_periods"], it["notes"],
            ])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].number_format = DATEF
        row[6].number_format = MONEY
        if isinstance(row[3].value, int) and row[3].value > 90:
            row[3].fill = WARN_FILL
        if row[8].value and isinstance(row[8].value, int) and row[8].value >= 3:
            row[7].font, row[7].fill = BAD_FONT, BAD_FILL
    if n == 0:
        ws.cell(row=2, column=6, value="No reconciling items submitted.")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(ws.max_row, 2)}"
    widths(ws, {1: 10, 2: 26, 3: 12, 4: 9, 5: 10, 6: 42, 7: 14, 8: 10, 9: 10, 10: 34})


def sheet_aging(wb, tb, asof: date):
    ws = wb.create_sheet("Aging")
    buckets = ("0-30", "31-60", "61-90", "91-180", "180+")
    heads = ["Account", "Name"] + list(buckets) + ["Oldest (days)", "Flag"]
    ws.append(heads)
    hdr(ws, len(heads))
    for a in sorted(tb, key=lambda x: x["account"]):
        if not a["items"]:
            continue
        by_bucket = defaultdict(lambda: ZERO)
        for it in a["items"]:
            by_bucket[it["bucket"]] += it["amount"]
        oldest = max(it["age_days"] for it in a["items"])
        flags = []
        if oldest > 180:
            flags.append("open item over 6 months - rarely timing at this age")
        if any(it["consecutive_periods"] >= 3 for it in a["items"]):
            flags.append("same item unresolved 3+ consecutive periods")
        ws.append([a["account"], a["name"]] +
                  [float(by_bucket.get(b, ZERO)) for b in buckets] +
                  [oldest, "; ".join(flags)])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col in range(3, 8):
            row[col - 1].number_format = MONEY
        if row[8].value:
            row[8].fill = WARN_FILL
    if ws.max_row == 1:
        ws.cell(row=2, column=2, value="No open reconciling items to age.")
    ws.freeze_panes = "A2"
    widths(ws, {1: 10, 2: 26, 3: 12, 4: 12, 5: 12, 6: 12, 7: 12, 8: 12, 9: 50})


def sheet_memo(wb, tb, meta, materiality, complete: bool, warnings: list[str]):
    ws = wb.create_sheet("Exceptions & Memo")
    widths(ws, {1: 4, 2: 100})
    r = 1

    def para(text, bold=False, size=11):
        nonlocal r
        c = ws.cell(row=r, column=2, value=text)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    para("RECONCILIATION PACK MEMO", bold=True, size=14)
    r += 1
    para(f"Client: {meta.get('Client', '')}    Period: {meta.get('Period', '')}", bold=True)
    para("Status: " + ("COMPLETE" if complete else "INCOMPLETE - see below"), bold=True)
    r += 1

    exceptions = [a for a in tb if a["status"] == "EXCEPTION"]
    not_received = [a for a in tb if a["status"] == "NOT RECEIVED"]
    unsigned = [a for a in tb if a["status"] in ("TIED - NOT SIGNED", "PREPARED - NOT REVIEWED")]

    if exceptions:
        para(f"EXCEPTIONS - {len(exceptions)} account(s) with an unexplained "
            f"difference. Never plugged; each stays open until explained.",
            bold=True, size=12)
        for a in sorted(exceptions, key=lambda x: -abs(x["unexplained"])):
            tag = " *** ESCALATE ***" if a["escalate"] else ""
            para(f"- {a['account']} {a['name']}: unexplained {a['unexplained']:,.2f}"
                f"{tag}")
        r += 1

    if not_received:
        para(f"NOT RECEIVED - {len(not_received)} account(s) with no supporting "
            f"schedule submitted yet.", bold=True, size=12)
        for a in sorted(not_received, key=lambda x: -abs(x["gl_net"])):
            tag = " *** ESCALATE ***" if a["escalate"] else ""
            para(f"- {a['account']} {a['name']}: GL balance {a['gl_net']:,.2f} "
                f"- suggested source: {a['hint']}{tag}")
        r += 1

    if unsigned:
        para(f"TIES BUT NOT FULLY SIGNED - {len(unsigned)} account(s) where the math "
            f"proves out but a preparer or reviewer signature is still missing.",
            bold=True, size=12)
        for a in unsigned:
            para(f"- {a['account']} {a['name']}: {a['status']}")
        r += 1

    recurring = [it for a in tb for it in a["items"] if it["consecutive_periods"] >= 3]
    if recurring:
        para(f"RECURRING ITEMS - {len(recurring)} item(s) unresolved for 3 or more "
            f"consecutive periods. This is almost never timing; treat as an open "
            f"finding, not an accepted balance.", bold=True, size=12)
        for it in recurring:
            para(f"- {it['account']} {it['description']} {it['amount']:,.2f} - "
                f"{it['consecutive_periods']} consecutive periods")
        r += 1

    if warnings:
        para("DATA WARNINGS", bold=True, size=12)
        for w in warnings:
            para("- " + w)
        r += 1

    para("Reminder: never plug. An account is reconciled only when the math ties "
        "to the cent AND a preparer and reviewer are both on record - not before.",
        bold=True)


# ------------------------------------------------------------------ main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--trial-balance", required=True)
    ap.add_argument("--schedules")
    ap.add_argument("--items")
    ap.add_argument("--prior-pack")
    ap.add_argument("--materiality")
    ap.add_argument("--as-of", help="pack as-of date YYYY-MM-DD, for aging")
    ap.add_argument("--client", default="")
    ap.add_argument("--period", default="")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    tb = load_tb(Path(args.trial_balance))
    schedules = load_schedules(Path(args.schedules) if args.schedules else None)
    items = load_items(Path(args.items) if args.items else None)
    prior = load_prior(Path(args.prior_pack) if args.prior_pack else None)
    materiality = dec(args.materiality) if args.materiality else None
    asof = pdate(args.as_of) if args.as_of else date.today()

    warnings = build_pack(tb, schedules, items, prior, materiality, asof)

    complete = all(a["status"] == "RECONCILED" for a in tb)

    meta = {
        "Client": args.client or "(not supplied)",
        "Period": args.period or "(not supplied)",
        "As-of date": asof.isoformat(),
        "Balance sheet accounts": len(tb),
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    wb = Workbook()
    sheet_summary(wb, tb, meta, materiality, complete)
    sheet_chart(wb, tb)
    sheet_items(wb, tb, asof)
    sheet_aging(wb, tb, asof)
    sheet_memo(wb, tb, meta, materiality, complete, warnings)
    wb.active = 0

    out = Path(args.out)
    if not complete:
        out = out.with_name(out.stem + " [INCOMPLETE]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    counts = defaultdict(int)
    for a in tb:
        counts[a["status"]] += 1

    print("=" * 70)
    print("BALANCE SHEET RECONCILIATION PACK")
    print("=" * 70)
    print(f"Balance sheet accounts: {len(tb)}")
    for s in STATUS_ORDER:
        print(f"  {s:<24} {counts.get(s, 0)}")
    print()
    print("COMPLETE" if complete else "INCOMPLETE - not every account is reconciled and signed")

    exceptions = [a for a in tb if a["status"] == "EXCEPTION"]
    if exceptions:
        print(f"\n{len(exceptions)} EXCEPTION(S) - unexplained difference, never plugged:")
        for a in exceptions:
            print(f"  ! {a['account']} {a['name']}: {a['unexplained']:,.2f}")
    not_received = [a for a in tb if a["status"] == "NOT RECEIVED"]
    if not_received:
        print(f"\n{len(not_received)} account(s) NOT RECEIVED (no schedule submitted).")
    if warnings:
        print("\nDATA WARNINGS:")
        for w in warnings:
            print(f"  ! {w}")
    print(f"\nWorkpaper: {out}")
    return 0 if complete else 1


if __name__ == "__main__":
    sys.exit(main())

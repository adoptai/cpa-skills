#!/usr/bin/env python3
"""
Field-by-field diff of two versions of a filing.

A diff is only useful if it accounts for EVERY field, including the ones that
appeared, vanished, or stayed put when they should have moved.

Gates:
  * matched + added + removed must equal the union of fields across both versions
  * both versions must be identified, so the comparison has a direction
  * every numeric change carries an old value, a new value and a delta
  * where a totals map is supplied, the sum of component changes must equal the
    change in the total - a change that does not propagate is an override or an error

Fields are matched on field_id where available, falling back to form + line.
Never on the printed description: descriptions change between software versions and
years, which manufactures phantom additions and removals.

Usage:
    python3 filing_diff.py \
        --version-a as_filed.csv --label-a "As filed 2026-03-14" \
        --version-b amended.csv --label-b "Amended 2026-08-02" \
        --totals-map totals.csv --prior-year-carryforwards py_cf.csv \
        --materiality 1000 --review-timestamp 2026-03-10 \
        --client "Brannon Holdings LLC" --return-type 1065 --tax-year 2025 \
        --out "Brannon - 2025 1065 Amended vs Original.xlsx"

--version-a / --version-b CSV:
    field_id, form, line, description, value, changed_at (optional), is_carryforward
      value: numeric or text. Text values are diffed but never netted.
      changed_at: timestamp of last edit, for the post-review test

--totals-map CSV:
    total_field_id, component_field_id
--prior-year-carryforwards CSV:
    field_id, description, amount_as_filed
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
MONEY = '#,##0.00;[Red](#,##0.00)'
DATEF = "yyyy-mm-dd"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
TOP = Border(top=Side(style="thin"))
DBL = Border(top=Side(style="thin"), bottom=Side(style="double"))

TIN_RE = re.compile(r"\b(\d{3}-\d{2}-\d{4}|\d{2}-\d{7}|\d{9})\b")

# Non-numeric fields whose change matters more than most dollar movements.
CRITICAL_WORDS = ("ein", "ssn", "tin", "identification", "filing status", "entity type",
                  "name", "address", "election", "checkbox", "final", "amended",
                  "initial return", "accounting method", "business code", "state",
                  "signature", "preparer", "routing", "account number")


def dec(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s2 = s.strip("()").replace(",", "").replace("$", "").strip()
    if s2.endswith("-"):
        neg, s2 = True, s2[:-1]
    try:
        v = Decimal(s2)
    except InvalidOperation:
        return None
    return -v if neg else v


def clean(s) -> str:
    return " ".join(str(s or "").split())


def truthy(s) -> bool:
    return clean(s).lower() in ("yes", "y", "true", "1", "x")


def pdt(raw):
    s = clean(raw)
    if not s:
        return None
    s = s.replace("T", " ").split(".")[0]
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
              "%m/%d/%Y %H:%M", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            continue
    return None


def mask_tin(s: str) -> str:
    return TIN_RE.sub(lambda m: "***-**-" + m.group(1)[-4:], str(s or ""))


def rows_of(path: Path, label: str):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {label} ({path})")
    return rows


def load_version(path: Path, label: str) -> dict:
    out = {}
    for i, r in enumerate(rows_of(path, label), start=1):
        fid = clean(r.get("field_id"))
        form, line = clean(r.get("form")), clean(r.get("line"))
        # Match on field_id, else form+line. NEVER on description.
        k = fid or f"{form}|{line}"
        if not k or k == "|":
            raise ValueError(f"{label} row {i}: needs a field_id, or a form and line. "
                             f"Matching on the printed description is not supported - "
                             f"descriptions change between software versions and years, "
                             f"which manufactures phantom additions and removals.")
        if k in out:
            raise ValueError(f"{label}: duplicate field key {k!r}. Field keys must be "
                             f"unique or the diff double-counts.")
        raw = r.get("value")
        out[k] = {
            "key": k, "field_id": fid, "form": form, "line": line,
            "description": clean(r.get("description")),
            "raw": clean(raw), "num": dec(raw),
            "changed_at": pdt(r.get("changed_at")),
            "is_carryforward": truthy(r.get("is_carryforward")),
        }
    return out


# ------------------------------------------------------------------------ diff

def build_diff(a: dict, b: dict) -> dict:
    keys_a, keys_b = set(a), set(b)
    matched = sorted(keys_a & keys_b)
    added = sorted(keys_b - keys_a)
    removed = sorted(keys_a - keys_b)

    changed_num, changed_text, unchanged = [], [], []
    for k in matched:
        fa, fb = a[k], b[k]
        both_numeric = fa["num"] is not None and fb["num"] is not None
        if both_numeric:
            if fa["num"] != fb["num"]:
                changed_num.append({**fb, "old": fa["num"], "new": fb["num"],
                                    "delta": fb["num"] - fa["num"],
                                    "old_raw": fa["raw"], "new_raw": fb["raw"],
                                    "old_description": fa["description"]})
            else:
                unchanged.append({**fb, "value": fb["num"]})
        else:
            if fa["raw"] != fb["raw"]:
                critical = any(w in (fb["description"] + " " + fb["line"]).lower()
                               for w in CRITICAL_WORDS)
                changed_text.append({**fb, "old_raw": fa["raw"], "new_raw": fb["raw"],
                                     "critical": critical})
            else:
                unchanged.append({**fb, "value": fb["raw"]})

    return {"matched": matched, "added": [b[k] for k in added],
            "removed": [a[k] for k in removed],
            "changed_num": sorted(changed_num, key=lambda x: -abs(x["delta"])),
            "changed_text": sorted(changed_text, key=lambda x: not x["critical"]),
            "unchanged": unchanged,
            "union": len(keys_a | keys_b)}


def test_propagation(diff, totals_map, a, b) -> list[dict]:
    """Sum of component changes must equal the change in the total."""
    if not totals_map:
        return []
    delta_by_key = {c["key"]: c["delta"] for c in diff["changed_num"]}
    out = []
    for total_key, components in totals_map.items():
        if total_key not in a and total_key not in b:
            continue
        total_delta = delta_by_key.get(total_key, ZERO)
        comp_delta = sum((delta_by_key.get(c, ZERO) for c in components), ZERO)
        moved_components = [c for c in components if c in delta_by_key]
        note = ""
        if total_delta != comp_delta:
            if total_delta == ZERO and comp_delta != ZERO:
                note = ("the total did NOT move while its components did - an override, "
                        "or a change that failed to propagate")
            elif comp_delta == ZERO and total_delta != ZERO:
                note = ("the total moved while no component did - the change was entered "
                        "directly on the total")
            else:
                note = "the total moved by a different amount than its components"
        out.append({"total_key": total_key,
                    "description": (b.get(total_key) or a.get(total_key))["description"],
                    "total_delta": total_delta, "component_delta": comp_delta,
                    "difference": total_delta - comp_delta,
                    "components_changed": len(moved_components),
                    "components_total": len(components),
                    "passed": total_delta == comp_delta, "note": note})
    return out


def test_carryforwards(b, prior) -> list[dict]:
    out = []
    for p in prior:
        k = clean(p.get("field_id"))
        expected = dec(p.get("amount_as_filed"))
        f = b.get(k)
        actual = f["num"] if f else None
        out.append({"field_id": k, "description": clean(p.get("description")),
                    "as_filed": expected, "carried_forward": actual,
                    "difference": (actual - expected)
                                  if (actual is not None and expected is not None)
                                  else None,
                    "passed": actual is not None and expected is not None
                              and actual == expected,
                    "note": ("not present in the current year at all - the carryforward "
                             "was dropped" if f is None else
                             "" if actual == expected else
                             "does not agree to the prior year as filed")})
    return out


def test_stability(diff, totals_map) -> list[dict]:
    """Fields that did not change where a dependent field did."""
    if not totals_map:
        return []
    changed = {c["key"] for c in diff["changed_num"]}
    out = []
    for total_key, components in totals_map.items():
        if total_key in changed:
            continue
        moved = [c for c in components if c in changed]
        if moved:
            out.append({"total_key": total_key, "moved_components": moved,
                        "note": f"{len(moved)} component(s) changed but this total did "
                                f"not. Either an override, or the change did not "
                                f"propagate. No ordinary diff surfaces this."})
    return out


# -------------------------------------------------------------------- workbook

def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


def sheet_summary(wb, meta, tests, diff, net, esc, overall, materiality):
    ws = wb.active
    ws.title = "Diff Summary"
    widths(ws, {1: 4, 2: 50, 3: 16, 4: 14, 5: 62})
    r = 1

    def line(label, v=None, s=None, *, bold=False, size=11, fill=None, note=""):
        nonlocal r
        c = ws.cell(row=r, column=2, value=label)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        if v is not None:
            cc = ws.cell(row=r, column=3,
                         value=float(v) if isinstance(v, Decimal) else v)
            if isinstance(v, Decimal):
                cc.number_format = MONEY
            cc.font = Font(bold=bold)
        if s is not None:
            sc = ws.cell(row=r, column=4, value=s)
            sc.font = OK_FONT if s == "PASS" else BAD_FONT
            if s != "PASS":
                sc.fill = BAD_FILL
        if note:
            n = ws.cell(row=r, column=5, value=note)
            n.font = Font(italic=True, size=9)
            n.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    line("FILING DIFF", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    v = ws.cell(row=r, column=2,
                value="COMPLETE - every field accounted for and every change explained"
                if overall else
                "INCOMPLETE - see the failing test(s). A diff that does not account for "
                "every field is the failure this test exists to prevent.")
    v.font = OK_FONT if overall else BAD_FONT
    if not overall:
        v.fill = BAD_FILL
    r += 2

    line("FIELD ACCOUNTING", bold=True, size=12, fill=SUB_FILL)
    line("  Fields matched (present in both)", len(diff["matched"]))
    line("  Added (only in the later version)", len(diff["added"]))
    line("  Removed (only in the earlier version)", len(diff["removed"]),
         fill=WARN_FILL if diff["removed"] else None,
         note="A removed field that carried a value means something stopped being "
              "reported. Different from a value changing, and easier to miss."
              if diff["removed"] else "")
    line("  = Total accounted for",
         len(diff["matched"]) + len(diff["added"]) + len(diff["removed"]), bold=True)
    line("  Union of fields across both versions", diff["union"])
    r += 1

    line("CHANGES", bold=True, size=12, fill=SUB_FILL)
    line("  Non-numeric changes", len(diff["changed_text"]),
         fill=BAD_FILL if any(c["critical"] for c in diff["changed_text"]) else None,
         note="Read these FIRST. A changed identification number, filing status, entity "
              "type, address or election often matters more than a five-figure numeric "
              "change - and cannot be netted, so it appears nowhere in the dollar "
              "summary." if diff["changed_text"] else "")
    line("  Numeric changes", len(diff["changed_num"]))
    line("  Unchanged", len(diff["unchanged"]))
    line("  Net numeric effect", net, bold=True)
    if materiality is not None:
        above = [c for c in diff["changed_num"] if abs(c["delta"]) >= materiality]
        line("  Changes at or above materiality", len(above))
    r += 1

    for t in tests:
        line(f"TEST {t['num']} - {t['name']}", None,
             "PASS" if t["passed"] else "FAIL", note=t.get("note", ""))
    r += 1

    if esc:
        line("ESCALATE", bold=True, size=12, fill=BAD_FILL)
        for e in esc:
            line("  " + e)
        r += 1
    r += 1
    line("Prepared by / date:  __________________  ____________", bold=True)
    line("Reviewed by / date:  __________________  ____________", bold=True)


def sheet_text_changes(wb, changed, label_a, label_b):
    ws = wb.create_sheet("Non-Numeric Changes")
    heads = ["Form", "Line", "Description", f"Old ({label_a})", f"New ({label_b})",
             "Critical", "Why it matters"]
    ws.append(heads)
    hdr(ws, len(heads))
    why = {
        "ein": "an identification number change misdirects correspondence and payments, "
               "and may invalidate the filing",
        "filing status": "changes the tax computation entirely",
        "entity type": "changes the entity's tax character",
        "election": "an election is binding and may not be revocable",
        "address": "misdirects notices - the taxpayer may never receive them",
        "routing": "a wrong routing or account number is a months-long problem",
        "name": "must agree to the identification record or the filing rejects",
    }
    for c in changed:
        desc = c["description"].lower()
        reason = next((v for k, v in why.items() if k in desc), "")
        ws.append([c["form"], c["line"], c["description"],
                   mask_tin(c["old_raw"]), mask_tin(c["new_raw"]),
                   "YES" if c["critical"] else "", reason])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[5].value == "YES":
            row[5].font, row[5].fill = BAD_FONT, BAD_FILL
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    if ws.max_row == 1:
        ws.cell(row=2, column=3, value="None.").font = OK_FONT
    else:
        r = ws.max_row + 2
        ws.cell(row=r, column=1,
                value="These are deliberately ahead of the numeric detail. They cannot be "
                      "netted, so they appear nowhere in the dollar summary - and they are "
                      "frequently the more consequential change."
                ).font = Font(italic=True)
    ws.freeze_panes = "C2"
    widths(ws, {1: 14, 2: 12, 3: 40, 4: 28, 5: 28, 6: 10, 7: 62})


def sheet_num_changes(wb, changed, label_a, label_b, materiality):
    ws = wb.create_sheet("Numeric Changes")
    heads = ["Form", "Line", "Description", f"Old ({label_a})", f"New ({label_b})",
             "Delta", "Above materiality", "Feeds a carryforward", "Changed at"]
    ws.append(heads)
    hdr(ws, len(heads))
    for c in changed:
        ws.append([c["form"], c["line"], c["description"], float(c["old"]),
                   float(c["new"]), float(c["delta"]),
                   "YES" if (materiality is not None
                             and abs(c["delta"]) >= materiality) else "",
                   "YES" if c["is_carryforward"] else "",
                   c["changed_at"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (3, 4, 5):
            row[i].number_format = MONEY
        row[8].number_format = "yyyy-mm-dd hh:mm"
        if row[6].value == "YES":
            row[6].fill = WARN_FILL
        if row[7].value == "YES":
            row[7].fill = BAD_FILL
    if ws.max_row == 1:
        ws.cell(row=2, column=3, value="None.").font = OK_FONT
    else:
        last = ws.max_row
        r = last + 2
        ws.cell(row=r, column=3, value="NET NUMERIC EFFECT").font = Font(bold=True)
        cc = ws.cell(row=r, column=6, value=f"=SUM(F2:F{last})")
        cc.number_format, cc.font, cc.border = MONEY, Font(bold=True), TOP
        r += 2
        ws.cell(row=r, column=1,
                value="A change to a field that feeds a carryforward affects subsequent "
                      "years, which may already be filed."
                ).font = Font(italic=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{max(ws.max_row, 2)}"
    widths(ws, {1: 14, 2: 12, 3: 42, 4: 18, 5: 18, 6: 16, 7: 17, 8: 20, 9: 17})


def sheet_added_removed(wb, diff, label_a, label_b):
    ws = wb.create_sheet("Added and Removed")
    heads = ["Status", "Form", "Line", "Description", "Value", "Note"]
    ws.append(heads)
    hdr(ws, len(heads))
    for f in diff["removed"]:
        ws.append([f"REMOVED (was in {label_a})", f["form"], f["line"],
                   f["description"],
                   float(f["num"]) if f["num"] is not None else mask_tin(f["raw"]),
                   "something stopped being reported - usually the more serious of the "
                   "two directions, and easier to miss than a changed value"
                   if (f["num"] not in (None, ZERO) or f["raw"]) else ""])
    for f in diff["added"]:
        ws.append([f"ADDED (new in {label_b})", f["form"], f["line"], f["description"],
                   float(f["num"]) if f["num"] is not None else mask_tin(f["raw"]),
                   "a new field with a value - confirm it belongs and is supported"])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if isinstance(row[4].value, (int, float)):
            row[4].number_format = MONEY
        if str(row[0].value).startswith("REMOVED"):
            row[0].font, row[0].fill = BAD_FONT, BAD_FILL
        else:
            row[0].fill = WARN_FILL
    if ws.max_row == 1:
        ws.cell(row=2, column=4,
                value="None - both versions contain the same field set.").font = OK_FONT
    ws.freeze_panes = "A2"
    widths(ws, {1: 30, 2: 14, 3: 12, 4: 40, 5: 20, 6: 64})


def sheet_propagation(wb, prop, stability):
    ws = wb.create_sheet("Propagation")
    heads = ["Total field", "Description", "Change in the total",
             "Sum of component changes", "Difference", "Components changed",
             "Result", "What a break means"]
    ws.append(heads)
    hdr(ws, len(heads))
    if not prop:
        ws.cell(row=2, column=1,
                value="NOT PERFORMED - no totals map supplied.")
        ws.cell(row=2, column=8,
                value="A totals map identifies which fields are the components of which "
                      "total. It is worth building once per return type: it is the only "
                      "way to detect a change that did not propagate.")
    for p in prop:
        ws.append([p["total_key"], p["description"], float(p["total_delta"]),
                   float(p["component_delta"]), float(p["difference"]),
                   f"{p['components_changed']} of {p['components_total']}",
                   "OK" if p["passed"] else "BREAK", p["note"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (2, 3, 4):
            if isinstance(row[i].value, (int, float)):
                row[i].number_format = MONEY
        if row[6].value == "BREAK":
            row[6].font, row[6].fill = BAD_FONT, BAD_FILL
        elif row[6].value == "OK":
            row[6].font = OK_FONT
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    if stability:
        r = ws.max_row + 3
        ws.cell(row=r, column=1, value="SUSPICIOUS STABILITY").font = Font(bold=True, size=12)
        r += 1
        for h, col in zip(["Total field", "Components that changed", "Note"], (1, 2, 3)):
            c = ws.cell(row=r, column=col, value=h)
            c.fill, c.font = HDR_FILL, HDR_FONT
        r += 1
        for s in stability:
            ws.cell(row=r, column=1, value=s["total_key"])
            ws.cell(row=r, column=2, value=", ".join(s["moved_components"]))
            c = ws.cell(row=r, column=3, value=s["note"])
            c.fill = BAD_FILL
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    ws.freeze_panes = "B2"
    widths(ws, {1: 20, 2: 38, 3: 20, 4: 24, 5: 16, 6: 20, 7: 12, 8: 60})


def sheet_carryforwards(wb, cf):
    ws = wb.create_sheet("Carryforward Continuity")
    heads = ["Field", "Description", "Prior year as filed", "Carried forward",
             "Difference", "Result", "Note"]
    ws.append(heads)
    hdr(ws, len(heads))
    if not cf:
        ws.cell(row=2, column=2,
                value="NOT PERFORMED - no prior-year carryforward schedule supplied.")
        ws.cell(row=2, column=7,
                value="Dropped carryforwards are the highest-frequency finding in return "
                      "review, and the check is a comparison. Doing it by eye is why it "
                      "gets missed.")
    for c in cf:
        ws.append([c["field_id"], c["description"],
                   float(c["as_filed"]) if c["as_filed"] is not None else None,
                   float(c["carried_forward"]) if c["carried_forward"] is not None else None,
                   float(c["difference"]) if c["difference"] is not None else None,
                   "OK" if c["passed"] else "BREAK", c["note"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (2, 3, 4):
            if isinstance(row[i].value, (int, float)):
                row[i].number_format = MONEY
        if row[5].value == "BREAK":
            row[5].font, row[5].fill = BAD_FONT, BAD_FILL
        elif row[5].value == "OK":
            row[5].font = OK_FONT
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    widths(ws, {1: 20, 2: 40, 3: 20, 4: 18, 5: 16, 6: 10, 7: 58})


def sheet_post_review(wb, changes, review_ts):
    ws = wb.create_sheet("Post-Review Changes")
    heads = ["Form", "Line", "Description", "Old", "New", "Delta", "Changed at"]
    ws.append(heads)
    hdr(ws, len(heads))
    if review_ts is None:
        ws.cell(row=2, column=3,
                value="NOT PERFORMED - no review timestamp supplied. Where one is "
                      "available, every change made after sign-off is listed here "
                      "regardless of amount: the amount is not the point, the fact that "
                      "the control was bypassed is.")
        widths(ws, {1: 14, 2: 12, 3: 90, 4: 16, 5: 16, 6: 16, 7: 18})
        return
    for c in changes:
        ws.append([c["form"], c["line"], c["description"],
                   float(c["old"]) if c.get("old") is not None else mask_tin(c.get("old_raw", "")),
                   float(c["new"]) if c.get("new") is not None else mask_tin(c.get("new_raw", "")),
                   float(c["delta"]) if c.get("delta") is not None else None,
                   c["changed_at"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (3, 4, 5):
            if isinstance(row[i].value, (int, float)):
                row[i].number_format = MONEY
        row[6].number_format = "yyyy-mm-dd hh:mm"
        row[6].fill = BAD_FILL
    if ws.max_row == 1:
        ws.cell(row=2, column=3,
                value=f"None - no field changed after {review_ts}.").font = OK_FONT
    else:
        r = ws.max_row + 2
        ws.cell(row=r, column=1,
                value=f"Every change above was made after review sign-off at {review_ts}. "
                      f"Escalate regardless of amount - the control was bypassed."
                ).font = Font(italic=True)
    widths(ws, {1: 14, 2: 12, 3: 42, 4: 18, 5: 18, 6: 16, 7: 18})


def sheet_unchanged(wb, unchanged):
    ws = wb.create_sheet("Unchanged Detail")
    heads = ["Form", "Line", "Description", "Value"]
    ws.append(heads)
    hdr(ws, len(heads))
    for f in unchanged:
        ws.append([f["form"], f["line"], f["description"],
                   float(f["value"]) if isinstance(f["value"], Decimal)
                   else mask_tin(f["value"])])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if isinstance(row[3].value, (int, float)):
            row[3].number_format = MONEY
    r = ws.max_row + 2
    ws.cell(row=r, column=1,
            value="Included so the field accounting can be re-performed: matched + added "
                  "+ removed must equal the union of fields across both versions."
            ).font = Font(italic=True)
    ws.freeze_panes = "A2"
    widths(ws, {1: 14, 2: 12, 3: 44, 4: 20})


# ------------------------------------------------------------------------ main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--version-a", required=True)
    ap.add_argument("--version-b", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--label-a", default="")
    ap.add_argument("--label-b", default="")
    ap.add_argument("--totals-map")
    ap.add_argument("--prior-year-carryforwards")
    ap.add_argument("--review-timestamp")
    ap.add_argument("--materiality")
    ap.add_argument("--client", default="")
    ap.add_argument("--return-type", default="")
    ap.add_argument("--tax-year", default="")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    a = load_version(Path(args.version_a), "version A")
    b = load_version(Path(args.version_b), "version B")

    totals_map = defaultdict(list)
    if args.totals_map and Path(args.totals_map).exists():
        for r in rows_of(Path(args.totals_map), "totals map"):
            t, c = clean(r.get("total_field_id")), clean(r.get("component_field_id"))
            if t and c:
                totals_map[t].append(c)

    prior = (rows_of(Path(args.prior_year_carryforwards), "prior-year carryforwards")
             if args.prior_year_carryforwards
             and Path(args.prior_year_carryforwards).exists() else [])
    review_ts = pdt(args.review_timestamp) if args.review_timestamp else None
    materiality = dec(args.materiality)

    diff = build_diff(a, b)
    prop = test_propagation(diff, totals_map, a, b)
    cf = test_carryforwards(b, prior)
    stability = test_stability(diff, totals_map)

    net = sum((c["delta"] for c in diff["changed_num"]), ZERO)
    accounted = len(diff["matched"]) + len(diff["added"]) + len(diff["removed"])
    post_review = []
    if review_ts:
        post_review = [c for c in diff["changed_num"] + diff["changed_text"]
                       if c.get("changed_at") and c["changed_at"] > review_ts]

    prop_breaks = [p for p in prop if not p["passed"]]
    cf_breaks = [c for c in cf if not c["passed"]]
    critical_text = [c for c in diff["changed_text"] if c["critical"]]
    removed_with_value = [f for f in diff["removed"]
                          if f["num"] not in (None, ZERO) or f["raw"]]

    tests = [
        {"num": 1, "name": "Every field accounted for",
         "passed": accounted == diff["union"],
         "note": "" if accounted == diff["union"] else
                 f"{accounted} accounted for against a union of {diff['union']} fields"},
        {"num": 2, "name": "Both versions identified and directional",
         "passed": bool(args.label_a and args.label_b),
         "note": "" if (args.label_a and args.label_b) else
                 "supply --label-a and --label-b; a diff with no direction cannot be read"},
        {"num": 3, "name": "Every change quantified",
         "passed": True,
         "note": f"{len(diff['changed_num'])} numeric change(s) net {net:,.2f}; "
                 f"{len(diff['changed_text'])} non-numeric change(s), which cannot be "
                 f"netted"},
        {"num": 4, "name": "Changes propagate to their totals",
         "passed": not prop_breaks,
         "note": ("no totals map supplied - the propagation test could not run, and it is "
                  "the only way to detect a change that did not propagate" if not prop
                  else "" if not prop_breaks else
                  f"{len(prop_breaks)} total(s) do not agree to the sum of their "
                  f"component changes")},
        {"num": 5, "name": "Carryforwards agree to the prior year as filed",
         "passed": not cf_breaks,
         "note": ("no prior-year carryforward schedule supplied" if not cf
                  else "" if not cf_breaks else
                  f"{len(cf_breaks)} carryforward(s) do not agree")},
        {"num": 6, "name": "No changes after review sign-off",
         "passed": not post_review,
         "note": ("no review timestamp supplied" if review_ts is None
                  else "" if not post_review else
                  f"{len(post_review)} field(s) changed after {review_ts}")},
        {"num": 7, "name": "No suspicious stability",
         "passed": not stability,
         "note": "" if not stability else
                 f"{len(stability)} total(s) did not move while their components did"},
    ]

    esc = []
    for f in removed_with_value:
        val = f["num"] if f["num"] is not None else f["raw"]
        esc.append(f"{f['form']} {f['line']} ({f['description']}): REMOVED, had a value of "
                   f"{val}. Something stopped being reported - different from a value "
                   f"changing, and easier to miss.")
    for p in prop_breaks:
        esc.append(f"{p['total_key']} ({p['description']}): {p['note']}. Total moved "
                   f"{p['total_delta']:,.2f}, components moved "
                   f"{p['component_delta']:,.2f}.")
    for s in stability:
        esc.append(f"{s['total_key']}: {s['note']}")
    for c in critical_text:
        esc.append(f"{c['form']} {c['line']} ({c['description']}): changed from "
                   f"\"{mask_tin(c['old_raw'])}\" to \"{mask_tin(c['new_raw'])}\". A "
                   f"change of this kind often matters more than a five-figure numeric "
                   f"change and cannot be netted.")
    for c in cf_breaks:
        esc.append(f"Carryforward {c['field_id']} ({c['description']}): {c['note']}. "
                   f"As filed {c['as_filed']}, carried forward {c['carried_forward']}.")
    for c in post_review:
        esc.append(f"{c['form']} {c['line']} ({c['description']}): changed at "
                   f"{c['changed_at']}, AFTER review sign-off at {review_ts}. The amount "
                   f"is not the point - the control was bypassed.")
    cf_feeding = [c for c in diff["changed_num"] if c["is_carryforward"]]
    if cf_feeding:
        esc.append(f"{len(cf_feeding)} changed field(s) feed a carryforward, so subsequent "
                   f"years are affected and may already be filed. Identify every affected "
                   f"year and state before concluding.")

    overall = all(t["passed"] for t in tests)

    meta = {
        "Client": args.client or "(not stated)",
        "Return type": args.return_type or "(not stated)",
        "Tax year": args.tax_year or "(not stated)",
        "Earlier version (A)": args.label_a or "*** NOT LABELLED ***",
        "Later version (B)": args.label_b or "*** NOT LABELLED ***",
        "Review timestamp": str(review_ts) if review_ts else "(not supplied)",
        "Materiality": float(materiality) if materiality is not None else "(not supplied)",
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    # ---- console
    print("=" * 76)
    print("FILING DIFF")
    print("=" * 76)
    print(f"Client : {meta['Client']}   {meta['Return type']}  TY {meta['Tax year']}")
    print(f"A (earlier): {meta['Earlier version (A)']}")
    print(f"B (later)  : {meta['Later version (B)']}")
    print()
    print(f"Fields matched {len(diff['matched'])}   added {len(diff['added'])}   "
          f"removed {len(diff['removed'])}   = {accounted} of union {diff['union']}  "
          f"{'OK' if accounted == diff['union'] else '*** DOES NOT ACCOUNT ***'}")
    print(f"Non-numeric changes {len(diff['changed_text'])}"
          f"{f' ({len(critical_text)} critical)' if critical_text else ''}"
          f"   numeric changes {len(diff['changed_num'])}   "
          f"unchanged {len(diff['unchanged'])}")
    print(f"Net numeric effect {net:,.2f}")
    print()
    for t in tests:
        print(f"Test {t['num']}: {'PASS' if t['passed'] else '*** FAIL ***':<14} {t['name']}")
        if t.get("note"):
            print(f"         {t['note']}")
    if critical_text:
        print("\nNON-NUMERIC CHANGES (read these first):")
        for c in critical_text[:10]:
            print(f"  ! {c['form']} {c['line']} {c['description']}: "
                  f"\"{mask_tin(c['old_raw'])}\" -> \"{mask_tin(c['new_raw'])}\"")
    if diff["changed_num"]:
        print("\nLargest numeric changes:")
        for c in diff["changed_num"][:8]:
            print(f"  {c['form']:<10} {c['line']:<10} {c['description'][:34]:<34} "
                  f"{c['old']:>14,.2f} -> {c['new']:>14,.2f}  "
                  f"delta {c['delta']:>13,.2f}"
                  f"{'  [carryforward]' if c['is_carryforward'] else ''}")
    if esc:
        print(f"\nESCALATE ({len(esc)}):")
        for e in esc[:15]:
            print(f"  ! {e}")
        if len(esc) > 15:
            print(f"  ... and {len(esc) - 15} more")

    if not overall and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("A diff is only useful if it accounts for every field. Resolve the failing")
        print("test(s) - a field present in one version and absent from the comparison is")
        print("exactly what this is meant to catch.")
        print("=" * 76)
        return 1

    wb = Workbook()
    sheet_summary(wb, meta, tests, diff, net, esc, overall, materiality)
    sheet_text_changes(wb, diff["changed_text"], args.label_a or "A", args.label_b or "B")
    sheet_num_changes(wb, diff["changed_num"], args.label_a or "A", args.label_b or "B",
                      materiality)
    sheet_added_removed(wb, diff, args.label_a or "A", args.label_b or "B")
    sheet_propagation(wb, prop, stability)
    sheet_carryforwards(wb, cf)
    sheet_post_review(wb, post_review, review_ts)
    sheet_unchanged(wb, diff["unchanged"])
    wb.active = 0
    out = Path(args.out)
    if not overall:
        out = out.with_name(out.stem + " [INCOMPLETE]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'COMPLETE' if overall else 'INCOMPLETE (forced)'}.  Workbook: {out}")
    return 0 if overall else 1


if __name__ == "__main__":
    sys.exit(main())

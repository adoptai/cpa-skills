#!/usr/bin/env python3
"""
Cutoff testing and search for unrecorded liabilities.

This procedure works on what was NOT recorded, so it cannot sample from its own
population. Gates reflect that:

  * the subsequent-disbursement population must tie to the cash account total for
    the search window - an incomplete population biases the result toward finding
    nothing, which is the direction that matters
  * every disbursement must be classified pre / post period end, and UNCLASSIFIED
    is a reported finding rather than a default
  * every pre-period-end item must be traced to recorded AP, an accrual, or a
    proposed adjustment
  * the accrual completeness checklist must be answered item by item, including
    "not applicable" with a reason

Classification depends on the SERVICE DATE - when goods transferred or services
were performed - not the invoice date and not the payment date. A missing service
date is treated as unclassified rather than guessed, because guessing here produces
a wrong answer that looks right.

Usage:
    python3 cutoff_test.py --disbursements subsequent_payments.csv \
        --disbursement-total 4182005.44 --recorded-ap recorded_ap.csv \
        --accruals accruals.csv --grni grni.csv --credit-memos credit_memos.csv \
        --period-end 2025-12-31 --search-end 2026-02-28 --materiality 25000 \
        --client "Ardenway Manufacturing" \
        --out "Ardenway - FY2025 Unrecorded Liabilities Search.xlsx"

--disbursements CSV:
    payment_date, payee, amount, invoice_ref, invoice_date, service_date,
    description, gl_account
--recorded-ap CSV:   vendor, invoice_ref, amount
--accruals CSV:      category, description, amount
--grni CSV:          po_no, vendor, item, receipt_date, value
--credit-memos CSV:  customer, invoice_no, memo_date, amount, reason
--checklist CSV:     category, answer, amount, basis      (optional; blanks reported)
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
MONEY = '#,##0.00;[Red](#,##0.00)'
DATEF = "yyyy-mm-dd"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
TOP = Border(top=Side(style="thin"))

ACCRUAL_CATEGORIES = [
    "Legal and professional fees",
    "Audit and tax fees for the period under audit",
    "Bonuses and incentive compensation",
    "Commissions",
    "Payroll and payroll taxes for the stub period",
    "Vacation and paid time off",
    "Utilities",
    "Rent and common area charges",
    "Property and other non-income taxes",
    "Interest",
    "Insurance",
    "Warranty and product returns",
    "Customer rebates and volume discounts",
    "Freight in transit",
    "Repairs performed but not invoiced",
    "Severance and restructuring",
    "Self-insured health claims incurred but not reported",
]


def dec(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def d0(raw):
    v = dec(raw)
    return v if v is not None else ZERO


def clean(s) -> str:
    return " ".join(str(s or "").split())


def key(s) -> str:
    return "".join(ch for ch in str(s or "").lower() if ch.isalnum())


def pdate(raw):
    s = clean(raw)
    if not s:
        return None
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s.split(" ")[0], f).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date {raw!r}. Prefer ISO YYYY-MM-DD.")


def rows_of(path: Path, label: str):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {label} ({path})")
    return rows


def load_disb(path: Path) -> list[dict]:
    out = []
    for i, r in enumerate(rows_of(path, "disbursements"), start=1):
        out.append({
            "row": i,
            "payment_date": pdate(r.get("payment_date")),
            "payee": clean(r.get("payee")) or "(no payee)",
            "amount": d0(r.get("amount")),
            "invoice_ref": clean(r.get("invoice_ref")),
            "invoice_date": pdate(r.get("invoice_date")),
            "service_date": pdate(r.get("service_date")),
            "description": clean(r.get("description")),
            "gl_account": clean(r.get("gl_account")),
            "classification": "", "traced_to": "", "adjustment": ZERO, "notes": [],
        })
    return out


def load_simple(path: Path | None, label: str) -> list[dict]:
    if not path or not path.exists():
        return []
    return rows_of(path, label)


# ----------------------------------------------------------------------- tests

def classify(disb, period_end: date, recorded_ap, accruals) -> None:
    ap_refs = {key(r.get("invoice_ref")) for r in recorded_ap if clean(r.get("invoice_ref"))}
    ap_vendors = defaultdict(lambda: ZERO)
    for r in recorded_ap:
        ap_vendors[key(r.get("vendor"))] += d0(r.get("amount"))
    accrual_total = sum((d0(r.get("amount")) for r in accruals), ZERO)
    accrual_words = " ".join(clean(r.get("description")).lower() +
                             " " + clean(r.get("category")).lower()
                             for r in accruals)

    for x in disb:
        sd = x["service_date"]
        if sd is None:
            x["classification"] = "unclassified"
            x["notes"].append(
                "no service date supplied - classification depends on when goods "
                "transferred or services were performed, not the invoice or payment "
                "date. Guessing here produces a wrong answer that looks right.")
            continue
        if sd <= period_end:
            x["classification"] = "belongs to the period"
            ref = key(x["invoice_ref"])
            if ref and ref in ap_refs:
                x["traced_to"] = "recorded AP"
            elif accrual_words and any(
                    w in accrual_words for w in
                    [w for w in clean(x["description"]).lower().split() if len(w) > 4]):
                x["traced_to"] = "possibly covered by an accrual - confirm"
                x["notes"].append("description overlaps an accrual description; confirm "
                                  "it is actually included rather than assuming")
            else:
                x["traced_to"] = "NOT RECORDED"
                x["adjustment"] = x["amount"]
                x["notes"].append(
                    f"service date {sd} is on or before period end {period_end} and the "
                    f"invoice is not in recorded AP - proposed liability")
        else:
            x["classification"] = "does not belong to the period"
            x["notes"].append(f"service date {sd} is after period end {period_end}")


def build_checklist(path: Path | None) -> list[dict]:
    supplied = {}
    if path and path.exists():
        for r in rows_of(path, "checklist"):
            supplied[clean(r.get("category")).lower()] = {
                "answer": clean(r.get("answer")),
                "amount": dec(r.get("amount")),
                "basis": clean(r.get("basis")),
            }
    out = []
    for cat in ACCRUAL_CATEGORIES:
        s = supplied.get(cat.lower(), {})
        answered = bool(s.get("answer"))
        reasoned = bool(s.get("basis"))
        out.append({
            "category": cat,
            "answer": s.get("answer", ""),
            "amount": s.get("amount"),
            "basis": s.get("basis", ""),
            "complete": answered and reasoned,
        })
    return out


# -------------------------------------------------------------------- workbook

def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


def sheet_summary(wb, meta, tests, stats, checklist, esc, overall, materiality):
    ws = wb.active
    ws.title = "Summary"
    widths(ws, {1: 4, 2: 52, 3: 14, 4: 18, 5: 60})
    r = 1

    def line(label, a=None, b=None, *, bold=False, size=11, fill=None, note=""):
        nonlocal r
        c = ws.cell(row=r, column=2, value=label)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        if a is not None:
            ws.cell(row=r, column=3, value=a).font = Font(bold=bold)
        if b is not None:
            cc = ws.cell(row=r, column=4, value=float(b))
            cc.number_format, cc.font = MONEY, Font(bold=bold)
        if note:
            n = ws.cell(row=r, column=5, value=note)
            n.font = Font(italic=True, size=9)
            n.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    line("SEARCH FOR UNRECORDED LIABILITIES AND CUTOFF TESTING", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    v = ws.cell(row=r, column=2,
                value="SEARCH COMPLETE - population proven and every item classified"
                if overall else
                "SEARCH INCOMPLETE - see the failing test(s). An incomplete population "
                "biases this procedure toward finding nothing.")
    v.font = OK_FONT if overall else BAD_FONT
    if not overall:
        v.fill = BAD_FILL
    r += 2

    line("UNRECORDED LIABILITIES IDENTIFIED", bold=True, size=12, fill=SUB_FILL)
    line("  From subsequent disbursements", stats["adj_n"], stats["adj_value"], bold=True)
    line("  From goods received not invoiced", stats["grni_n"], stats["grni_value"])
    line("  From unprocessed invoices", stats["unproc_n"], stats["unproc_value"])
    line("  TOTAL PROPOSED ADJUSTMENT", None, stats["total_adjustment"], bold=True)
    if materiality is not None:
        line("  Materiality", None, materiality)
        over = stats["total_adjustment"] > materiality
        line("  Exceeds materiality?", "YES" if over else "no", bold=True,
             fill=BAD_FILL if over else None,
             note="This changes the financial statements." if over else "")
    r += 1

    ws.cell(row=r, column=3, value="Count").font = Font(bold=True)
    ws.cell(row=r, column=4, value="Value").font = Font(bold=True)
    r += 1
    line("CLASSIFICATION OF SUBSEQUENT DISBURSEMENTS", bold=True, size=12, fill=SUB_FILL)
    for k in ("belongs to the period", "does not belong to the period", "unclassified"):
        b = stats["by_class"][k]
        line("  " + k, b["n"], b["value"],
             fill=BAD_FILL if k == "unclassified" and b["n"] else None,
             note="Unclassified is a finding, not a default - no service date was "
                  "supplied, so these could not be tested."
                  if k == "unclassified" and b["n"] else "")
    r += 1

    for t in tests:
        line(f"TEST {t['num']} - {t['name']}", "PASS" if t["passed"] else "FAIL",
             bold=True, fill=None if t["passed"] else BAD_FILL, note=t.get("note", ""))
        c = ws.cell(row=r - 1, column=3)
        c.font = OK_FONT if t["passed"] else BAD_FONT
    r += 1

    incomplete = [c for c in checklist if not c["complete"]]
    line("ACCRUAL COMPLETENESS CHECKLIST", bold=True, size=12, fill=SUB_FILL)
    line("  Categories", len(checklist))
    line("  Answered with a basis", len(checklist) - len(incomplete))
    line("  Unanswered or without a basis", len(incomplete),
         fill=BAD_FILL if incomplete else None,
         note="'Not applicable' is a valid answer with a reason. A blank is not."
              if incomplete else "")
    r += 1

    if esc:
        line("ESCALATE", bold=True, size=12, fill=BAD_FILL)
        for e in esc:
            line("  " + e)
        r += 1
    r += 1
    line("Prepared by / date:  __________________  ____________", bold=True)
    line("Reviewed by / date:  __________________  ____________", bold=True)


def sheet_proof(wb, stats, disb_total, listed_total, period_end, search_end):
    ws = wb.create_sheet("Population Proof")
    widths(ws, {1: 4, 2: 50, 3: 14, 4: 18, 5: 62})
    r = 1
    ws.cell(row=r, column=2, value="SEARCH POPULATION PROOF").font = Font(bold=True, size=14)
    r += 2
    ws.cell(row=r, column=2, value=f"Search window: {period_end} to {search_end}"
            ).font = Font(bold=True)
    r += 2
    if disb_total is None:
        ws.cell(row=r, column=2,
                value="Cash account total for the search window NOT SUPPLIED. In this "
                      "procedure specifically, an incomplete population biases the result "
                      "toward finding nothing - which is the direction that matters. "
                      "Obtain the total.").fill = BAD_FILL
        r += 2
    else:
        for label, v in [("Disbursements listed", listed_total),
                         ("Total per the cash account for the window", disb_total),
                         ("Difference", listed_total - disb_total)]:
            ws.cell(row=r, column=2, value=label).font = Font(bold=label == "Difference")
            c = ws.cell(row=r, column=4, value=float(v))
            c.number_format = MONEY
            if label == "Difference":
                c.font = OK_FONT if v == ZERO else BAD_FONT
                if v != ZERO:
                    c.fill = BAD_FILL
            r += 1
        r += 1
    ws.cell(row=r, column=2, value="Classification balance").font = Font(bold=True, size=12)
    r += 1
    ws.cell(row=r, column=3, value="Count").font = Font(bold=True)
    ws.cell(row=r, column=4, value="Value").font = Font(bold=True)
    r += 1
    tot_n = tot_v = 0
    for k in ("belongs to the period", "does not belong to the period", "unclassified"):
        b = stats["by_class"][k]
        ws.cell(row=r, column=2, value="  " + k)
        ws.cell(row=r, column=3, value=b["n"])
        c = ws.cell(row=r, column=4, value=float(b["value"]))
        c.number_format = MONEY
        tot_n += b["n"]
        tot_v += b["value"]
        r += 1
    ws.cell(row=r, column=2, value="= Total classified").font = Font(bold=True)
    c1 = ws.cell(row=r, column=3, value=tot_n)
    c2 = ws.cell(row=r, column=4, value=float(tot_v))
    c2.number_format = MONEY
    for c in (c1, c2):
        c.font, c.border = Font(bold=True), TOP
    r += 1
    ws.cell(row=r, column=2, value="Population")
    ws.cell(row=r, column=3, value=stats["n"])
    c = ws.cell(row=r, column=4, value=float(stats["value"]))
    c.number_format = MONEY
    r += 1
    ws.cell(row=r, column=2, value="Difference").font = Font(bold=True)
    d1 = ws.cell(row=r, column=3, value=tot_n - stats["n"])
    d2 = ws.cell(row=r, column=4, value=float(tot_v - stats["value"]))
    d2.number_format = MONEY
    ok = tot_n == stats["n"] and tot_v == stats["value"]
    for c in (d1, d2):
        c.font = OK_FONT if ok else BAD_FONT
        if not ok:
            c.fill = BAD_FILL


def sheet_disb(wb, disb):
    ws = wb.create_sheet("Subsequent Disbursements")
    heads = ["Payment date", "Payee", "Amount", "Invoice ref", "Invoice date",
             "Service date", "Description", "GL account", "Classification",
             "Traced to", "Proposed adjustment", "Notes"]
    ws.append(heads)
    hdr(ws, len(heads))
    order = {"belongs to the period": 0, "unclassified": 1,
             "does not belong to the period": 2}
    for x in sorted(disb, key=lambda y: (order.get(y["classification"], 3),
                                         -y["amount"])):
        ws.append([x["payment_date"], x["payee"], float(x["amount"]),
                   x["invoice_ref"] or None, x["invoice_date"], x["service_date"],
                   x["description"] or None, x["gl_account"] or None,
                   x["classification"], x["traced_to"] or None,
                   float(x["adjustment"]) if x["adjustment"] else None,
                   "; ".join(x["notes"])])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (0, 4, 5):
            row[i].number_format = DATEF
        row[2].number_format = MONEY
        row[10].number_format = MONEY
        if row[8].value == "unclassified":
            row[8].font, row[8].fill = BAD_FONT, BAD_FILL
        elif row[8].value == "belongs to the period":
            row[8].fill = WARN_FILL
        if row[9].value == "NOT RECORDED":
            row[9].font, row[9].fill = BAD_FONT, BAD_FILL
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{max(ws.max_row, 2)}"
    widths(ws, {1: 13, 2: 28, 3: 15, 4: 16, 5: 12, 6: 12, 7: 34, 8: 14, 9: 26,
                10: 30, 11: 18, 12: 66})


def sheet_adjustments(wb, disb, grni, unproc, materiality):
    ws = wb.create_sheet("Proposed Adjustments")
    heads = ["#", "Source", "Description", "Debit (expense/asset)", "Credit (liability)",
             "Reference", "Basis"]
    ws.append(heads)
    hdr(ws, len(heads))
    n = 0
    for x in disb:
        if x["adjustment"] == ZERO:
            continue
        n += 1
        ws.append([n, "subsequent disbursement",
                   f"{x['payee']} - {x['description'][:44]}",
                   float(x["adjustment"]), None,
                   x["invoice_ref"] or x["payment_date"],
                   f"service date {x['service_date']} on or before period end, not in "
                   f"recorded AP"])
        ws.append(["", "", "  Accounts payable / accrued liabilities", None,
                   float(x["adjustment"]), "", ""])
    for g in grni:
        v = d0(g.get("value"))
        if v == ZERO:
            continue
        n += 1
        ws.append([n, "goods received not invoiced",
                   f"{clean(g.get('vendor'))} - {clean(g.get('item'))[:40]}",
                   float(v), None, clean(g.get("po_no")),
                   f"received {clean(g.get('receipt_date'))} with no invoice"])
        ws.append(["", "", "  Accrued liabilities", None, float(v), "", ""])
    for u in unproc:
        v = d0(u.get("amount"))
        if v == ZERO:
            continue
        n += 1
        ws.append([n, "unprocessed invoice",
                   f"{clean(u.get('vendor'))} - {clean(u.get('description'))[:40]}",
                   float(v), None, clean(u.get("invoice_ref")),
                   "invoice received before the books closed and not recorded"])
        ws.append(["", "", "  Accounts payable", None, float(v), "", ""])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[3].number_format = MONEY
        row[4].number_format = MONEY
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    if n == 0:
        ws.cell(row=2, column=3, value="None - no unrecorded liabilities identified.")
    else:
        last = ws.max_row
        r = last + 2
        ws.cell(row=r, column=3, value="TOTAL (must be equal)").font = Font(bold=True)
        for col, letter in ((4, "D"), (5, "E")):
            c = ws.cell(row=r, column=col, value=f"=SUM({letter}2:{letter}{last})")
            c.number_format, c.font, c.border = MONEY, Font(bold=True), TOP
        if materiality is not None:
            r += 1
            ws.cell(row=r, column=3, value="Materiality").font = Font(bold=True)
            c = ws.cell(row=r, column=5, value=float(materiality))
            c.number_format = MONEY
    widths(ws, {1: 5, 2: 28, 3: 46, 4: 20, 5: 20, 6: 18, 7: 58})


def sheet_grni(wb, grni, unproc, period_end):
    ws = wb.create_sheet("GRNI and Unprocessed")
    ws.cell(row=1, column=1, value="GOODS RECEIVED NOT INVOICED").font = Font(bold=True, size=12)
    heads = ["PO", "Vendor", "Item", "Receipt date", "Value", "Accrual required"]
    ws.append(heads)
    hdr(ws, len(heads), row=2)
    for g in grni:
        rd = pdate(g.get("receipt_date"))
        ws.append([clean(g.get("po_no")), clean(g.get("vendor")), clean(g.get("item")),
                   rd, float(d0(g.get("value"))),
                   "YES - received before period end" if rd and rd <= period_end
                   else "no - received after period end"])
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        row[3].number_format = DATEF
        row[4].number_format = MONEY
        if "YES" in str(row[5].value):
            row[5].fill = WARN_FILL
    if ws.max_row == 2:
        ws.cell(row=3, column=1, value="None supplied.")
    r = ws.max_row + 3
    ws.cell(row=r, column=1, value="UNPROCESSED INVOICES").font = Font(bold=True, size=12)
    r += 1
    for h, col in zip(["Vendor", "Invoice ref", "Invoice date", "Amount", "Description"],
                      range(1, 6)):
        c = ws.cell(row=r, column=col, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    r += 1
    if not unproc:
        ws.cell(row=r, column=1, value="None supplied. If the client maintains no "
                                       "unprocessed invoice file, say so - it is a "
                                       "scope point as well as a control one.")
    for u in unproc:
        ws.cell(row=r, column=1, value=clean(u.get("vendor")))
        ws.cell(row=r, column=2, value=clean(u.get("invoice_ref")))
        d = pdate(u.get("invoice_date"))
        c = ws.cell(row=r, column=3, value=d)
        c.number_format = DATEF
        c2 = ws.cell(row=r, column=4, value=float(d0(u.get("amount"))))
        c2.number_format = MONEY
        ws.cell(row=r, column=5, value=clean(u.get("description")))
        r += 1
    widths(ws, {1: 22, 2: 18, 3: 13, 4: 16, 5: 40, 6: 34})


def sheet_checklist(wb, checklist):
    ws = wb.create_sheet("Accrual Checklist")
    heads = ["Category", "Answer", "Amount accrued", "Basis", "Status"]
    ws.append(heads)
    hdr(ws, len(heads))
    for c in checklist:
        ws.append([c["category"], c["answer"] or None,
                   float(c["amount"]) if c["amount"] is not None else None,
                   c["basis"] or None,
                   "complete" if c["complete"] else "NOT ANSWERED"])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].number_format = MONEY
        if row[4].value != "complete":
            row[4].font, row[4].fill = BAD_FONT, BAD_FILL
        else:
            row[4].font = OK_FONT
    r = ws.max_row + 2
    for note in [
        "'Not applicable' is a valid answer with a reason. A blank is not.",
        "Audit and tax fees for the period under audit deserve their own mention: they are",
        "incurred, the amount is known to the firm performing the work, and they are",
        "omitted with remarkable frequency.",
    ]:
        ws.cell(row=r, column=1, value=note).font = Font(italic=True)
        r += 1
    ws.freeze_panes = "B2"
    widths(ws, {1: 46, 2: 18, 3: 18, 4: 56, 5: 16})


def sheet_memos(wb, memos, period_end):
    ws = wb.create_sheet("Credit Memos")
    heads = ["Customer", "Invoice", "Memo date", "Amount", "Reason", "Effect"]
    ws.append(heads)
    hdr(ws, len(heads))
    for m in memos:
        d = pdate(m.get("memo_date"))
        ws.append([clean(m.get("customer")), clean(m.get("invoice_no")), d,
                   float(d0(m.get("amount"))), clean(m.get("reason")),
                   "reverses period revenue - test the original sale"
                   if d and d > period_end else "within the period"])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].number_format = DATEF
        row[3].number_format = MONEY
        if "reverses" in str(row[5].value):
            row[5].fill = WARN_FILL
    if ws.max_row == 1:
        ws.cell(row=2, column=1,
                value="None supplied. Subsequent credit memos test the revenue and "
                      "receivable direction and are worth requesting.")
    widths(ws, {1: 28, 2: 16, 3: 12, 4: 16, 5: 40, 6: 46})


# ------------------------------------------------------------------------ main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--disbursements", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--disbursement-total")
    ap.add_argument("--recorded-ap")
    ap.add_argument("--accruals")
    ap.add_argument("--grni")
    ap.add_argument("--unprocessed")
    ap.add_argument("--credit-memos")
    ap.add_argument("--checklist")
    ap.add_argument("--period-end", required=True)
    ap.add_argument("--search-end", required=True)
    ap.add_argument("--materiality")
    ap.add_argument("--client", default="")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    period_end = pdate(args.period_end)
    search_end = pdate(args.search_end)
    disb = load_disb(Path(args.disbursements))
    recorded_ap = load_simple(Path(args.recorded_ap) if args.recorded_ap else None, "AP")
    accruals = load_simple(Path(args.accruals) if args.accruals else None, "accruals")
    grni = load_simple(Path(args.grni) if args.grni else None, "GRNI")
    unproc = load_simple(Path(args.unprocessed) if args.unprocessed else None,
                         "unprocessed invoices")
    memos = load_simple(Path(args.credit_memos) if args.credit_memos else None,
                        "credit memos")
    checklist = build_checklist(Path(args.checklist) if args.checklist else None)
    materiality = dec(args.materiality)
    disb_total = dec(args.disbursement_total)

    classify(disb, period_end, recorded_ap, accruals)

    listed_total = sum((x["amount"] for x in disb), ZERO)
    by_class = {k: {"n": 0, "value": ZERO} for k in
                ("belongs to the period", "does not belong to the period", "unclassified")}
    for x in disb:
        b = by_class[x["classification"]]
        b["n"] += 1
        b["value"] += x["amount"]

    adj_items = [x for x in disb if x["adjustment"] != ZERO]
    grni_pre = [g for g in grni
                if pdate(g.get("receipt_date")) and pdate(g.get("receipt_date")) <= period_end]
    unproc_pre = [u for u in unproc
                  if pdate(u.get("invoice_date")) and pdate(u.get("invoice_date")) <= period_end]
    grni_value = sum((d0(g.get("value")) for g in grni_pre), ZERO)
    unproc_value = sum((d0(u.get("amount")) for u in unproc_pre), ZERO)
    adj_value = sum((x["adjustment"] for x in adj_items), ZERO)

    stats = {"n": len(disb), "value": listed_total, "by_class": by_class,
             "adj_n": len(adj_items), "adj_value": adj_value,
             "grni_n": len(grni_pre), "grni_value": grni_value,
             "unproc_n": len(unproc_pre), "unproc_value": unproc_value,
             "total_adjustment": adj_value + grni_value + unproc_value}

    unclassified = by_class["unclassified"]
    incomplete = [c for c in checklist if not c["complete"]]
    memos_after = [m for m in memos
                   if pdate(m.get("memo_date")) and pdate(m.get("memo_date")) > period_end]

    tests = [
        {"num": 1, "name": "Search population ties to the cash account",
         "passed": disb_total is not None and listed_total == disb_total,
         "note": "" if disb_total is not None and listed_total == disb_total else
                 ("cash account total not supplied" if disb_total is None else
                  f"difference {listed_total - disb_total:,.2f} - an incomplete "
                  f"population biases this procedure toward finding nothing")},
        {"num": 2, "name": "Every disbursement classified",
         "passed": unclassified["n"] == 0,
         "note": "" if unclassified["n"] == 0 else
                 f"{unclassified['n']} item(s) totalling {unclassified['value']:,.2f} "
                 f"have no service date and could not be classified"},
        {"num": 3, "name": "Pre-period-end payments traced to AP, an accrual, or an adjustment",
         "passed": True,
         "note": f"{len(adj_items)} item(s) totalling {adj_value:,.2f} were not recorded "
                 f"and are proposed as adjustments"},
        {"num": 4, "name": "Goods received before period end are accrued",
         "passed": True,
         "note": f"{len(grni_pre)} GRNI item(s) totalling {grni_value:,.2f} require accrual"
                 if grni_pre else ("no GRNI population supplied" if not grni else "none")},
        {"num": 5, "name": "Unprocessed invoices identified",
         "passed": True,
         "note": f"{len(unproc_pre)} invoice(s) totalling {unproc_value:,.2f}"
                 if unproc_pre else
                 "no unprocessed invoice file supplied - if the client maintains none, "
                 "that is a scope point as well as a control one"},
        {"num": 6, "name": "Accrual completeness checklist answered",
         "passed": not incomplete,
         "note": "" if not incomplete else
                 f"{len(incomplete)} of {len(checklist)} categories unanswered or without "
                 f"a basis. 'Not applicable' is a valid answer with a reason; a blank is not"},
        {"num": 7, "name": "Subsequent credit memos reviewed",
         "passed": True,
         "note": f"{len(memos_after)} credit memo(s) after period end"
                 if memos else "none supplied - these test the revenue direction"},
    ]

    esc = []
    if disb_total is None:
        esc.append("Cash account total for the search window not supplied. In this "
                   "procedure an incomplete population biases the result toward finding "
                   "nothing, which is the direction that matters.")
    elif listed_total != disb_total:
        esc.append(f"The disbursement listing does not tie to the cash account by "
                   f"{listed_total - disb_total:,.2f}. Obtain the complete listing.")
    if unclassified["n"]:
        esc.append(f"{unclassified['n']} disbursement(s) totalling "
                   f"{unclassified['value']:,.2f} have no service date and are "
                   f"unclassified. These are a scope limitation, not a pass.")
    if stats["total_adjustment"] > ZERO:
        line = (f"Unrecorded liabilities of {stats['total_adjustment']:,.2f} identified")
        if materiality is not None:
            line += (f", against materiality of {materiality:,.2f}"
                     + (" - EXCEEDS MATERIALITY" if stats["total_adjustment"] > materiality
                        else ""))
        esc.append(line + ".")
    audit_fee = next((c for c in checklist
                      if "audit and tax fees" in c["category"].lower()), None)
    if audit_fee and not audit_fee["complete"]:
        esc.append("Audit and tax fees for the period under audit are unanswered on the "
                   "checklist. They are incurred, the amount is known to the firm doing "
                   "the work, and they are omitted with remarkable frequency.")
    held = [x for x in disb if x["invoice_date"] and x["service_date"]
            and x["service_date"] <= period_end and x["invoice_date"] > period_end]
    if len(held) >= 3:
        esc.append(f"{len(held)} invoice(s) for pre-period-end services are dated after "
                   f"period end. A pattern of invoices held and entered late is worth "
                   f"examining.")
    if memos_after:
        esc.append(f"{len(memos_after)} credit memo(s) issued after period end totalling "
                   f"{sum((d0(m.get('amount')) for m in memos_after), ZERO):,.2f} - test "
                   f"the original sales.")

    overall = all(t["passed"] for t in tests)

    meta = {
        "Client": args.client or "(not stated)",
        "Period end": str(period_end),
        "Search window end": str(search_end),
        "Disbursements examined": len(disb),
        "Materiality": float(materiality) if materiality is not None else "(not supplied)",
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    # ---- console
    print("=" * 76)
    print("SEARCH FOR UNRECORDED LIABILITIES AND CUTOFF TESTING")
    print("=" * 76)
    print(f"Client : {meta['Client']}   Period end {period_end}   "
          f"Search to {search_end}")
    print(f"Disbursements examined {len(disb):,}   value {listed_total:,.2f}")
    if disb_total is not None:
        d = listed_total - disb_total
        print(f"Cash account total    {disb_total:>16,.2f}   difference {d:>14,.2f}  "
              f"{'TIES' if d == ZERO else '*** DOES NOT TIE ***'}")
    else:
        print("Cash account total    NOT SUPPLIED")
    print()
    for k in ("belongs to the period", "does not belong to the period", "unclassified"):
        b = by_class[k]
        print(f"  {k:<32} {b['n']:>5}   {b['value']:>15,.2f}")
    print()
    print(f"UNRECORDED LIABILITIES IDENTIFIED")
    print(f"  from subsequent disbursements   {adj_value:>15,.2f}  ({len(adj_items)} items)")
    print(f"  from goods received not invoiced{grni_value:>15,.2f}  ({len(grni_pre)} items)")
    print(f"  from unprocessed invoices       {unproc_value:>15,.2f}  ({len(unproc_pre)} items)")
    print(f"  TOTAL PROPOSED ADJUSTMENT       {stats['total_adjustment']:>15,.2f}")
    if materiality is not None:
        print(f"  materiality                     {materiality:>15,.2f}   "
              f"{'*** EXCEEDS MATERIALITY ***' if stats['total_adjustment'] > materiality else 'below'}")
    print()
    for t in tests:
        print(f"Test {t['num']}: {'PASS' if t['passed'] else '*** FAIL ***':<14} {t['name']}")
        if t.get("note"):
            print(f"         {t['note']}")
    if esc:
        print("\nESCALATE:")
        for e in esc:
            print(f"  ! {e}")

    if not overall and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("This procedure works on what was NOT recorded, so an incomplete population")
        print("or an unclassified item is a real gap, not a rounding issue.")
        print("=" * 76)
        return 1

    wb = Workbook()
    sheet_summary(wb, meta, tests, stats, checklist, esc, overall, materiality)
    sheet_proof(wb, stats, disb_total, listed_total, period_end, search_end)
    sheet_disb(wb, disb)
    sheet_adjustments(wb, disb, grni_pre, unproc_pre, materiality)
    sheet_grni(wb, grni, unproc, period_end)
    sheet_checklist(wb, checklist)
    sheet_memos(wb, memos, period_end)
    wb.active = 0
    out = Path(args.out)
    if not overall:
        out = out.with_name(out.stem + " [INCOMPLETE]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'COMPLETE' if overall else 'INCOMPLETE (forced)'}.  Workbook: {out}")
    return 0 if overall else 1


if __name__ == "__main__":
    sys.exit(main())

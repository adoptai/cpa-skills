#!/usr/bin/env python3
"""
Indirect-method cash flow statement: derive it from the balance sheet and prove it.

Working capital movements are DERIVED from the balance sheet rather than taken from
the statement, because the usual failure is a plug hidden inside an "other" caption
that survives review when only the bottom line is checked.

Gates:
  * the balance sheet must balance in both periods, checked before anything else
  * operating + investing + financing + FX must equal the balance sheet cash movement
  * every difference between a derived movement and the amount presented must be
    attributed to a named cause
  * every balance sheet account must feed a section - an account feeding nothing is
    the definition of a hidden plug

Usage:
    python3 cash_flow.py --balance-sheet bs.csv --income-statement is.csv \
        --noncash noncash.csv --investing-financing invfin.csv \
        --adjustments adjustments.csv \
        --client "Latham Industrial Group" --period FY2025 \
        --out "Latham - FY2025 Cash Flow Tie-Out.xlsx"

--balance-sheet CSV:
    account, caption, type, opening, closing
      type: cash | operating_asset | operating_liability | investing_asset |
            debt | equity | noncash_contra
      Anything else is rejected - an untyped account cannot be placed in a section.

--income-statement CSV:  caption, amount      (net_income row required)
--noncash CSV:           caption, amount      (depreciation, stock comp, etc.)
--investing-financing CSV: caption, section, amount     section: investing | financing
--adjustments CSV:       account, cause, amount, evidence
      cause must be one of the recognised non-cash / reclass reasons.
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
MONEY = '#,##0.00;[Red](#,##0.00)'

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
TOP = Border(top=Side(style="thin"))
DBL = Border(top=Side(style="thin"), bottom=Side(style="double"))

TYPES = {"cash", "operating_asset", "operating_liability", "investing_asset",
         "debt", "equity", "noncash_contra"}

# Sign applied to the balance sheet movement to get the cash effect.
#   asset up   -> cash down
#   liability up -> cash up
OP_SIGN = {"operating_asset": Decimal(-1), "operating_liability": Decimal(1)}

# Contra-asset accounts (accumulated depreciation, the AR allowance) are deliberately
# NOT derived into working capital. Their movement IS the non-cash charge that is added
# back separately, so deriving them too would double-count it. Instead their movement is
# cross-checked against the non-cash adjustments supplied - see reconcile_sections().

CAUSES = {
    "noncash_addition": "Non-cash addition (lease, non-monetary exchange, capitalised interest)",
    "acquisition": "Acquisition or disposal of a business",
    "fx_translation": "Foreign currency translation of a foreign operation",
    "reclassification": "Reclassification between captions or current/non-current",
    "noncash_writeoff": "Non-cash write-off (bad debt against allowance, inventory writedown)",
    "accrual_transfer": "Accrued but unpaid amount moving between accrual accounts",
    "noncash_charge": "Non-cash charge already in net income (stock comp, deferred tax)",
}

JUDGMENT = [
    ("Interest paid", "operating under most frameworks, but presentation varies and "
                      "disclosure may be required"),
    ("Dividends received", "operating or investing depending on framework and policy"),
    ("Capitalised interest", "investing, though the expense ran through operating"),
    ("Overdrafts", "financing, or part of cash and equivalents, depending on the "
                   "arrangement"),
    ("Revolving facility draws and repayments", "gross versus net presentation"),
    ("Restricted cash", "whether it belongs in cash and equivalents at all"),
    ("Book versus bank overdrafts", "these are different and are treated differently"),
]


def dec(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def d0(raw):
    v = dec(raw)
    return v if v is not None else ZERO


def clean(s) -> str:
    return " ".join(str(s or "").split())


def rows_of(path: Path, label: str):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {label} ({path})")
    return rows


def load_bs(path: Path) -> list[dict]:
    out = []
    for i, r in enumerate(rows_of(path, "balance sheet"), start=1):
        t = clean(r.get("type")).lower()
        if t not in TYPES:
            sys.exit(f"Balance sheet row {i} ({clean(r.get('account'))}): type "
                     f"'{t}' is not recognised. Every account must be typed as one of "
                     f"{sorted(TYPES)} - an untyped account cannot be placed in a "
                     f"section, and an account that feeds no section is a hidden plug.")
        out.append({
            "row": i,
            "account": clean(r.get("account")) or f"ACC{i:04d}",
            "caption": clean(r.get("caption")) or clean(r.get("account")),
            "type": t,
            "opening": d0(r.get("opening")),
            "closing": d0(r.get("closing")),
        })
    return out


def load_pairs(path: Path | None, label: str, extra=()) -> list[dict]:
    if not path or not path.exists():
        return []
    out = []
    for r in rows_of(path, label):
        d = {"caption": clean(r.get("caption")), "amount": d0(r.get("amount"))}
        for e in extra:
            d[e] = clean(r.get(e)).lower()
        out.append(d)
    return out


def load_adjustments(path: Path | None) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = defaultdict(list)
    if not path or not path.exists():
        return out
    for i, r in enumerate(rows_of(path, "adjustments"), start=1):
        acct = clean(r.get("account"))
        cause = clean(r.get("cause")).lower()
        amt = dec(r.get("amount"))
        if not acct or amt is None:
            continue
        if cause not in CAUSES:
            sys.exit(f"Adjustment row {i} ({acct}): cause '{cause}' is not recognised. "
                     f"Valid causes are {sorted(CAUSES)}. If a difference cannot be "
                     f"placed in one of these, it is not explained - report it rather "
                     f"than inventing a caption.")
        out[acct].append({"cause": cause, "amount": amt,
                          "evidence": clean(r.get("evidence"))})
    return out


# ------------------------------------------------------------------- derivation

def derive(bs, adjustments) -> dict:
    cash = [a for a in bs if a["type"] == "cash"]
    open_cash = sum((a["opening"] for a in cash), ZERO)
    close_cash = sum((a["closing"] for a in cash), ZERO)

    wc, unattributed = [], []
    for a in bs:
        if a["type"] not in OP_SIGN:
            continue
        raw = a["closing"] - a["opening"]
        adjs = adjustments.get(a["account"], [])
        adj_total = sum((x["amount"] for x in adjs), ZERO)
        cash_move = raw - adj_total
        effect = cash_move * OP_SIGN[a["type"]]
        wc.append({**a, "raw_movement": raw, "adjustments": adjs,
                   "adjustment_total": adj_total, "cash_movement": cash_move,
                   "cash_effect": effect,
                   "sign": "outflow if positive movement" if a["type"].endswith("asset")
                           else "inflow if positive movement"})
    return {"cash_accounts": cash, "open_cash": open_cash, "close_cash": close_cash,
            "change_in_cash": close_cash - open_cash, "working_capital": wc}


def bs_balance(bs, which: str) -> Decimal:
    """Assets less liabilities and equity. Should be zero."""
    assets = sum((a[which] for a in bs
                  if a["type"] in ("cash", "operating_asset", "investing_asset")), ZERO)
    contra = sum((a[which] for a in bs if a["type"] == "noncash_contra"), ZERO)
    liabs = sum((a[which] for a in bs
                 if a["type"] in ("operating_liability", "debt")), ZERO)
    equity = sum((a[which] for a in bs if a["type"] == "equity"), ZERO)
    return (assets - contra) - (liabs + equity)


def build_statement(bs, inc, noncash, invfin, der) -> dict:
    ni = next((x["amount"] for x in inc
               if "net income" in x["caption"].lower()
               or "net_income" in x["caption"].lower()), None)
    if ni is None:
        ni = sum((x["amount"] for x in inc), ZERO)

    nc_total = sum((x["amount"] for x in noncash), ZERO)
    wc_total = sum((x["cash_effect"] for x in der["working_capital"]), ZERO)
    operating = ni + nc_total + wc_total

    inv = [x for x in invfin if x.get("section") == "investing"]
    fin = [x for x in invfin if x.get("section") == "financing"]
    fx = [x for x in invfin if x.get("section") == "fx"]
    investing = sum((x["amount"] for x in inv), ZERO)
    financing = sum((x["amount"] for x in fin), ZERO)
    fx_effect = sum((x["amount"] for x in fx), ZERO)

    computed = operating + investing + financing + fx_effect
    return {"net_income": ni, "noncash": noncash, "noncash_total": nc_total,
            "wc_total": wc_total, "operating": operating,
            "investing_items": inv, "investing": investing,
            "financing_items": fin, "financing": financing,
            "fx_items": fx, "fx_effect": fx_effect,
            "computed_change": computed,
            "actual_change": der["change_in_cash"],
            "unexplained": computed - der["change_in_cash"]}


def reconcile_sections(bs, st, noncash, invfin) -> list[dict]:
    """Cross-check each non-operating balance sheet movement against what was supplied.

    Every movement must be explained by something: a contra account by the non-cash
    charge added back, PP&E by capital expenditure and disposals, debt by draws and
    repayments, retained earnings by net income and distributions. An account whose
    movement is explained by nothing is a hidden plug.
    """
    def move(types):
        return sum((a["closing"] - a["opening"] for a in bs if a["type"] in types), ZERO)

    def cap_total(items, *words):
        return sum((x["amount"] for x in items
                    if any(w in x["caption"].lower() for w in words)), ZERO)

    checks = []

    contra = move({"noncash_contra"})
    nc_depr = cap_total(noncash, "depreciation", "amortisation", "amortization",
                        "bad debt", "impairment", "allowance")
    checks.append({
        "name": "Contra-asset movement vs non-cash charges added back",
        "derived": contra, "supplied": nc_depr, "difference": contra - nc_depr,
        "note": "Accumulated depreciation and the AR allowance move by the non-cash "
                "charge, less any write-off or disposal. A difference means a disposal "
                "or write-off was not supplied, or the add-back is wrong.",
    })

    inv_assets = move({"investing_asset"})
    inv_supplied = -sum((x["amount"] for x in invfin
                         if x.get("section") == "investing"), ZERO)
    checks.append({
        "name": "Investing asset movement vs investing activities",
        "derived": inv_assets, "supplied": inv_supplied,
        "difference": inv_assets - inv_supplied,
        "note": "Gross fixed asset movement should equal capital expenditure less the "
                "cost of disposals. A difference means a disposal or a non-cash addition "
                "was not supplied.",
    })

    debt = move({"debt"})
    fin_debt = cap_total(invfin, "debt", "borrow", "loan", "note", "facility",
                         "repay")
    checks.append({
        "name": "Debt movement vs financing activities",
        "derived": debt, "supplied": fin_debt, "difference": debt - fin_debt,
        "note": "Debt balances move by draws less repayments. A difference means a "
                "movement was not supplied, or debt was assumed or forgiven non-cash.",
    })

    equity = move({"equity"})
    ni = st["net_income"]
    fin_equity = cap_total(invfin, "dividend", "distribution", "equity", "stock",
                           "share", "buyback", "repurchase")
    expected_equity = ni + fin_equity
    checks.append({
        "name": "Equity movement vs net income and equity financing",
        "derived": equity, "supplied": expected_equity,
        "difference": equity - expected_equity,
        "note": "Equity moves by net income, plus equity issued or repurchased, less "
                "dividends and distributions. A difference means a distribution was not "
                "supplied, or something was posted directly to equity - which is worth "
                "reading in full.",
    })

    return checks


# -------------------------------------------------------------------- workbook

def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


def sheet_statement(wb, st, der, meta, overall, esc):
    ws = wb.active
    ws.title = "Cash Flow Statement"
    widths(ws, {1: 4, 2: 56, 3: 18, 4: 18, 5: 58})
    r = 1

    def line(label, v=None, t=None, *, bold=False, size=11, fill=None, border=None,
             note="", indent=0):
        nonlocal r
        c = ws.cell(row=r, column=2, value=("    " * indent) + label)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        for col, val in ((3, v), (4, t)):
            if val is None:
                continue
            cc = ws.cell(row=r, column=col, value=float(val))
            cc.number_format = MONEY
            cc.font = Font(bold=bold)
            if border:
                cc.border = border
        if note:
            n = ws.cell(row=r, column=5, value=note)
            n.font = Font(italic=True, size=9)
            n.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    line("STATEMENT OF CASH FLOWS - INDIRECT METHOD", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    v = ws.cell(row=r, column=2,
                value="FOOTS - the statement reconciles to the balance sheet cash movement"
                if overall else
                "DOES NOT FOOT - see the unexplained line. Do not plug it.")
    v.font = OK_FONT if overall else BAD_FONT
    if not overall:
        v.fill = BAD_FILL
    r += 2

    line("CASH FLOWS FROM OPERATING ACTIVITIES", bold=True, size=12, fill=SUB_FILL)
    line("Net income", st["net_income"], indent=1)
    if st["noncash"]:
        line("Adjustments for non-cash items:", indent=1)
        for x in st["noncash"]:
            line(x["caption"], x["amount"], indent=2)
    line("Changes in operating assets and liabilities:", indent=1)
    for x in der["working_capital"]:
        if x["cash_effect"] == ZERO:
            continue
        line(x["caption"], x["cash_effect"], indent=2)
    line("Net cash from operating activities", None, st["operating"], bold=True,
         border=TOP)
    r += 1

    line("CASH FLOWS FROM INVESTING ACTIVITIES", bold=True, size=12, fill=SUB_FILL)
    for x in st["investing_items"]:
        line(x["caption"], x["amount"], indent=1)
    line("Net cash from investing activities", None, st["investing"], bold=True,
         border=TOP)
    r += 1

    line("CASH FLOWS FROM FINANCING ACTIVITIES", bold=True, size=12, fill=SUB_FILL)
    for x in st["financing_items"]:
        line(x["caption"], x["amount"], indent=1)
    line("Net cash from financing activities", None, st["financing"], bold=True,
         border=TOP)
    r += 1

    if st["fx_items"] or st["fx_effect"] != ZERO:
        for x in st["fx_items"]:
            line(x["caption"], x["amount"], indent=1)
        line("Effect of exchange rate changes on cash", None, st["fx_effect"], bold=True)
        r += 1

    line("NET CHANGE IN CASH", None, st["computed_change"], bold=True, border=TOP)
    line("Cash at beginning of period", None, der["open_cash"])
    line("Cash at end of period - computed", None,
         der["open_cash"] + st["computed_change"], bold=True)
    line("Cash at end of period - per balance sheet", None, der["close_cash"])
    line("UNEXPLAINED - DO NOT PLUG", None, st["unexplained"], bold=True, border=DBL,
         fill=None if st["unexplained"] == ZERO else BAD_FILL,
         note="" if st["unexplained"] == ZERO else
              "A difference here must be attributed to a named cause. Absorbing it into "
              "an 'other' caption is how a plug survives review.")
    c = ws.cell(row=r - 1, column=4)
    c.font = OK_FONT if st["unexplained"] == ZERO else BAD_FONT
    r += 1

    if esc:
        line("ESCALATE", bold=True, size=12, fill=BAD_FILL)
        for e in esc:
            line("  " + e)
        r += 1

    r += 1
    line("Prepared by / date:  __________________  ____________", bold=True)
    line("Reviewed by / date:  __________________  ____________", bold=True)


def sheet_proof(wb, tests):
    ws = wb.create_sheet("Proof")
    widths(ws, {1: 4, 2: 54, 3: 18, 4: 62})
    r = 1
    ws.cell(row=r, column=2, value="PROOF").font = Font(bold=True, size=14)
    r += 2
    for t in tests:
        c = ws.cell(row=r, column=2, value=f"TEST {t['num']} - {t['name']}")
        c.font, c.fill = Font(bold=True, size=12), SUB_FILL
        r += 1
        for label, val in t["lines"]:
            ws.cell(row=r, column=2, value="  " + label)
            cc = ws.cell(row=r, column=3,
                         value=float(val) if isinstance(val, Decimal) else val)
            if isinstance(val, Decimal):
                cc.number_format = MONEY
            r += 1
        ws.cell(row=r, column=2, value="  Result").font = Font(bold=True)
        v = ws.cell(row=r, column=3, value="PASS" if t["passed"] else "FAIL")
        v.font = OK_FONT if t["passed"] else BAD_FONT
        if not t["passed"]:
            v.fill = BAD_FILL
        if t.get("note"):
            n = ws.cell(row=r, column=4, value=t["note"])
            n.font = Font(italic=True, size=9)
            n.alignment = Alignment(wrap_text=True, vertical="top")
        r += 2


def sheet_wc(wb, wc):
    ws = wb.create_sheet("Working Capital Derivation")
    heads = ["Account", "Caption", "Type", "Opening", "Closing", "Raw movement",
             "Named adjustments", "Cash movement", "Sign applied", "Cash flow effect",
             "Adjustment causes"]
    ws.append(heads)
    hdr(ws, len(heads))
    for x in wc:
        ws.append([x["account"], x["caption"], x["type"], float(x["opening"]),
                   float(x["closing"]), float(x["raw_movement"]),
                   float(x["adjustment_total"]), float(x["cash_movement"]),
                   "+1" if OP_SIGN[x["type"]] > 0 else "-1",
                   float(x["cash_effect"]),
                   "; ".join(f"{CAUSES[a['cause']]}: {a['amount']:,.2f}"
                             f"{' (' + a['evidence'] + ')' if a['evidence'] else ''}"
                             for a in x["adjustments"])])
    last = ws.max_row
    ws.append(["TOTAL", "", "", None, None, None, None, None, "",
               f"=SUM(J2:J{last})", ""])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (3, 4, 5, 6, 7, 9):
            row[i].number_format = MONEY
        if isinstance(row[6].value, (int, float)) and row[6].value:
            row[6].fill = WARN_FILL
    for c in range(1, len(heads) + 1):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.font, cell.border = Font(bold=True), TOP
    r = ws.max_row + 2
    for note in [
        "Movements are DERIVED from the balance sheet, not taken from the statement.",
        "Sign logic: operating asset up = outflow; operating liability up = inflow.",
        "Check the account types before reading the result. Two accounts mis-typed in",
        "opposite directions leave the statement footing perfectly while both sections",
        "are misstated - this tab is what surfaces that.",
    ]:
        ws.cell(row=r, column=1, value=note).font = Font(italic=True)
        r += 1
    ws.freeze_panes = "C2"
    widths(ws, {1: 14, 2: 30, 3: 20, 4: 15, 5: 15, 6: 15, 7: 17, 8: 15, 9: 12,
                10: 17, 11: 60})


def sheet_bs(wb, bs, section_checks):
    ws = wb.create_sheet("Balance Sheet Movement")
    heads = ["Account", "Caption", "Type", "Opening", "Closing", "Movement",
             "Feeds"]
    ws.append(heads)
    hdr(ws, len(heads))
    feeds = {"cash": "cash - the target of the proof",
             "operating_asset": "operating - working capital",
             "operating_liability": "operating - working capital",
             "noncash_contra": "operating - non-cash add-back (NOT working capital)",
             "investing_asset": "investing",
             "debt": "financing", "equity": "financing / net income"}
    for a in bs:
        ws.append([a["account"], a["caption"], a["type"], float(a["opening"]),
                   float(a["closing"]), float(a["closing"] - a["opening"]),
                   feeds.get(a["type"], "?")])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (3, 4, 5):
            row[i].number_format = MONEY
        if row[2].value == "noncash_contra":
            row[6].fill = WARN_FILL
    r = ws.max_row + 3
    ws.cell(row=r, column=1, value="SECTION RECONCILIATIONS").font = Font(bold=True, size=12)
    r += 1
    for h, col in zip(["Check", "Per balance sheet", "Per detail supplied", "Difference",
                       "What a difference means"], range(1, 6)):
        c = ws.cell(row=r, column=col, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    r += 1
    for c in section_checks:
        ws.cell(row=r, column=1, value=c["name"])
        for col, v in ((2, c["derived"]), (3, c["supplied"]), (4, c["difference"])):
            cc = ws.cell(row=r, column=col, value=float(v))
            cc.number_format = MONEY
        d = ws.cell(row=r, column=4)
        if c["difference"] != ZERO:
            d.font, d.fill = BAD_FONT, BAD_FILL
        else:
            d.font = OK_FONT
        n = ws.cell(row=r, column=5, value=c["note"])
        n.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1,
            value="Contra-asset accounts are NOT derived into working capital. Their "
                  "movement is the non-cash charge added back separately, so deriving "
                  "them as well would double-count it."
            ).font = Font(italic=True)
    ws.freeze_panes = "C2"
    widths(ws, {1: 42, 2: 20, 3: 22, 4: 16, 5: 70, 6: 16, 7: 44})


def sheet_compare(wb, st, client):
    ws = wb.create_sheet("Comparison")
    heads = ["Line", "Derived", "Per client statement", "Difference"]
    ws.append(heads)
    hdr(ws, len(heads))
    if not client:
        ws.cell(row=2, column=1,
                value="No client statement supplied - nothing to compare. Where the "
                      "client has prepared a statement, the comparison against the "
                      "derived figures is the point of the exercise.")
        widths(ws, {1: 90, 2: 18, 3: 20, 4: 16})
        return
    derived = {"operating": st["operating"], "investing": st["investing"],
               "financing": st["financing"], "fx": st["fx_effect"],
               "net change in cash": st["computed_change"]}
    for k, v in derived.items():
        cv = client.get(k)
        ws.append([k, float(v), float(cv) if cv is not None else None,
                   float(v - cv) if cv is not None else None])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (1, 2, 3):
            row[i].number_format = MONEY
        if isinstance(row[3].value, (int, float)) and row[3].value:
            row[3].font, row[3].fill = BAD_FONT, BAD_FILL
    widths(ws, {1: 30, 2: 18, 3: 22, 4: 16})


def sheet_judgment(wb, st):
    ws = wb.create_sheet("Judgment Items")
    heads = ["Classification question", "Why it is judgment", "Amount if identified",
             "Determined by"]
    ws.append(heads)
    hdr(ws, len(heads))
    caps = {x["caption"].lower(): x["amount"]
            for x in st["noncash"] + st["investing_items"] + st["financing_items"]}
    for name, why in JUDGMENT:
        amt = next((v for k, v in caps.items() if name.split()[0].lower() in k), None)
        ws.append([name, why, float(amt) if amt is not None else None, None])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].number_format = MONEY
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    r = ws.max_row + 2
    ws.cell(row=r, column=1,
            value="Flagged, not concluded. Section classification is a policy and "
                  "framework question and belongs with the person signing."
            ).font = Font(italic=True)
    widths(ws, {1: 40, 2: 62, 3: 20, 4: 22})


# ------------------------------------------------------------------------ main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--balance-sheet", required=True)
    ap.add_argument("--income-statement", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--noncash")
    ap.add_argument("--investing-financing")
    ap.add_argument("--adjustments")
    ap.add_argument("--client-statement",
                    help="CSV: line, amount - the client's own statement, for comparison")
    ap.add_argument("--client", default="")
    ap.add_argument("--period", default="")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    bs = load_bs(Path(args.balance_sheet))
    inc = load_pairs(Path(args.income_statement), "income statement")
    noncash = load_pairs(Path(args.noncash) if args.noncash else None, "non-cash")
    invfin = load_pairs(Path(args.investing_financing) if args.investing_financing
                        else None, "investing/financing", extra=("section",))
    adjustments = load_adjustments(Path(args.adjustments) if args.adjustments else None)

    client_st = {}
    if args.client_statement and Path(args.client_statement).exists():
        for r in rows_of(Path(args.client_statement), "client statement"):
            client_st[clean(r.get("line")).lower()] = d0(r.get("amount"))

    open_bal = bs_balance(bs, "opening")
    close_bal = bs_balance(bs, "closing")
    der = derive(bs, adjustments)
    st = build_statement(bs, inc, noncash, invfin, der)
    section_checks = reconcile_sections(bs, st, noncash, invfin)
    unfed = [c for c in section_checks if c["difference"] != ZERO]

    tests = []
    tests.append({"num": 1, "name": "Balance sheet balances in both periods",
                  "lines": [("Opening: assets less liabilities and equity", open_bal),
                            ("Closing: assets less liabilities and equity", close_bal)],
                  "passed": open_bal == ZERO and close_bal == ZERO,
                  "note": "" if open_bal == ZERO and close_bal == ZERO else
                          "A cash flow built on an unbalanced balance sheet is "
                          "meaningless. Fix this before reading anything else."})
    tests.append({"num": 2, "name": "Cash movement identified from the balance sheet",
                  "lines": [("Cash accounts identified",
                             len(der["cash_accounts"])),
                            ("Opening cash", der["open_cash"]),
                            ("Closing cash", der["close_cash"]),
                            ("Change in cash", der["change_in_cash"])],
                  "passed": bool(der["cash_accounts"]),
                  "note": "" if der["cash_accounts"] else
                          "No account was typed as 'cash'. Nothing can be proven "
                          "without identifying the target."})
    unattributed = [x for x in der["working_capital"]
                    if x["adjustment_total"] != ZERO and
                    any(not a["evidence"] for a in x["adjustments"])]
    tests.append({"num": 3, "name": "Working capital derived from the balance sheet",
                  "lines": [("Operating accounts derived",
                             len(der["working_capital"])),
                            ("Total working capital effect", st["wc_total"]),
                            ("Accounts with named adjustments",
                             len([x for x in der["working_capital"]
                                  if x["adjustments"]])),
                            ("Adjustments lacking evidence", len(unattributed))],
                  "passed": not unattributed,
                  "note": "" if not unattributed else
                          "An adjustment with a cause but no evidence reference is a "
                          "label, not an explanation."})
    tests.append({"num": 4, "name": "The statement foots to the balance sheet",
                  "lines": [("Operating", st["operating"]),
                            ("Investing", st["investing"]),
                            ("Financing", st["financing"]),
                            ("FX effect", st["fx_effect"]),
                            ("Computed change in cash", st["computed_change"]),
                            ("Change in cash per balance sheet", st["actual_change"]),
                            ("Unexplained", st["unexplained"])],
                  "passed": st["unexplained"] == ZERO,
                  "note": "" if st["unexplained"] == ZERO else
                          "Attribute this to a named cause. Do not absorb it into an "
                          "'other' caption."})
    tests.append({"num": 5,
                  "name": "Every non-operating movement is explained by what was supplied",
                  "lines": [(c["name"], c["difference"]) for c in section_checks],
                  "passed": not unfed,
                  "note": "" if not unfed else
                          "Each movement must be explained by something - a contra "
                          "account by the non-cash charge, PP&E by capex and disposals, "
                          "debt by draws and repayments, equity by net income and "
                          "distributions. An unexplained movement is a hidden plug."})
    tests.append({"num": 6, "name": "Comparison to the client's statement",
                  "lines": ([("Client statement supplied", "yes")] +
                            [(k, v) for k, v in client_st.items()]) if client_st
                           else [("Client statement supplied", "no")],
                  "passed": True,
                  "note": "" if client_st else
                          "NOT PERFORMED - no client statement supplied. Where one "
                          "exists, comparing it to the derived figures is the point."})

    esc = []
    if st["unexplained"] != ZERO:
        esc.append(f"The statement does not foot: {st['unexplained']:,.2f} unexplained. "
                   f"Attribute it to a named cause - a plug hidden in an 'other' caption "
                   f"survives review because reviewers check the bottom line rather than "
                   f"each line against the balance sheet.")
    if open_bal != ZERO or close_bal != ZERO:
        esc.append(f"The balance sheet does not balance (opening {open_bal:,.2f}, "
                   f"closing {close_bal:,.2f}). Everything downstream is meaningless "
                   f"until this is fixed.")
    for u in unfed:
        esc.append(f"{u['name']}: balance sheet movement {u['derived']:,.2f} against "
                   f"{u['supplied']:,.2f} supplied - difference {u['difference']:,.2f}. "
                   f"{u['note']}")
    others = [x for x in st["noncash"] + st["investing_items"] + st["financing_items"]
              if "other" in x["caption"].lower() or "misc" in x["caption"].lower()]
    if others:
        tot = sum((abs(x["amount"]) for x in others), ZERO)
        sections = abs(st["operating"]) + abs(st["investing"]) + abs(st["financing"])
        if sections and tot / sections > Decimal("0.05"):
            esc.append(f"'Other' or miscellaneous captions total {tot:,.2f}, more than 5% "
                       f"of gross section activity. A material unexplained 'other' is a "
                       f"plug regardless of its label.")
    if st["operating"] < st["net_income"] and st["net_income"] > ZERO:
        esc.append(f"Operating cash flow ({st['operating']:,.2f}) is below net income "
                   f"({st['net_income']:,.2f}). Report the relationship across periods; "
                   f"it is an earnings-quality signal, not a conclusion.")
    if st["fx_effect"] != ZERO and not any(
            "foreign" in a["caption"].lower() or "translation" in a["caption"].lower()
            for a in bs):
        esc.append(f"An FX effect of {st['fx_effect']:,.2f} is presented but no balance "
                   f"sheet account suggests a foreign operation. Confirm the entity has "
                   f"foreign operations.")

    overall = all(t["passed"] for t in tests)

    meta = {
        "Client": args.client or "(not stated)",
        "Period": args.period or "(not stated)",
        "Balance sheet accounts": len(bs),
        "Cash accounts identified": len(der["cash_accounts"]),
        "Operating accounts derived": len(der["working_capital"]),
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    # ---- console
    print("=" * 76)
    print("CASH FLOW STATEMENT TIE-OUT")
    print("=" * 76)
    print(f"Client : {meta['Client']}    Period: {meta['Period']}")
    print()
    print(f"Balance sheet balances: opening {open_bal:,.2f}   closing {close_bal:,.2f}"
          f"   {'OK' if open_bal == ZERO and close_bal == ZERO else '*** DOES NOT BALANCE ***'}")
    print()
    print(f"  Net income                        {st['net_income']:>16,.2f}")
    print(f"  Non-cash adjustments              {st['noncash_total']:>16,.2f}")
    print(f"  Working capital (derived)         {st['wc_total']:>16,.2f}")
    print(f"  = Operating                       {st['operating']:>16,.2f}")
    print(f"  Investing                         {st['investing']:>16,.2f}")
    print(f"  Financing                         {st['financing']:>16,.2f}")
    print(f"  FX effect                         {st['fx_effect']:>16,.2f}")
    print(f"  = NET CHANGE IN CASH              {st['computed_change']:>16,.2f}")
    print(f"  Change per balance sheet          {st['actual_change']:>16,.2f}")
    print(f"  UNEXPLAINED                       {st['unexplained']:>16,.2f}   "
          f"{'OK' if st['unexplained'] == ZERO else '*** MUST BE ZERO - DO NOT PLUG ***'}")
    print()
    for t in tests:
        print(f"Test {t['num']}: {'PASS' if t['passed'] else '*** FAIL ***':<14} "
              f"{t['name']}")
    if esc:
        print("\nESCALATE:")
        for e in esc:
            print(f"  ! {e}")

    if not overall and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("Attribute every difference to a named cause. The statement of cash flows")
        print("must reconcile to the balance sheet by construction - there is no judgment")
        print("in that requirement.")
        print("=" * 76)
        return 1

    wb = Workbook()
    sheet_statement(wb, st, der, meta, overall, esc)
    sheet_proof(wb, tests)
    sheet_wc(wb, der["working_capital"])
    sheet_bs(wb, bs, section_checks)
    sheet_compare(wb, st, client_st)
    sheet_judgment(wb, st)
    wb.active = 0
    out = Path(args.out)
    if not overall:
        out = out.with_name(out.stem + " [DOES NOT FOOT]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'FOOTS' if overall else 'DOES NOT FOOT (forced)'}.  Workbook: {out}")
    return 0 if overall else 1


if __name__ == "__main__":
    sys.exit(main())

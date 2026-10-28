#!/usr/bin/env python3
"""
Work through an auditor's PBC (provided-by-client) request list against a
local document folder listing, and report which requests have candidate
evidence and which are gaps.

This is a pure local matcher: it never walks the filesystem itself (the user
supplies the folder listing, e.g. via `find . -type f` redirected to CSV) and
it never invents its own keywords from the free-text request description -
matching is keyword-driven and the keywords are supplied by the client or the
engagement team, because deciding what filename pattern "counts" as evidence
for a request is a judgment call this script does not make.

Gates:
  * every request resolves to exactly one status: matched, gap, or
    no_keywords_supplied - a request with a blank keywords column is never
    silently folded into "gap," because that would claim a document search
    was performed when it wasn't.
  * every matched status names the specific file(s) and the keyword that
    triggered each match, so a reviewer can verify it isn't a false positive.
  * every request_id in the input appears exactly once in the output - never
    dropped, never duplicated.
  * a match is reported as "candidate evidence located," never as "request
    satisfied" - whether the document actually fulfills the request remains
    the engagement team's determination.

Usage - emit blank templates:
    python3 pbc_assembler.py --template-requests pbc_requests.csv
    python3 pbc_assembler.py --template-folder document_folder.csv

Usage - run the assembler:
    python3 pbc_assembler.py --requests pbc_requests.csv \
        --folder document_folder.csv --prior-year prior_year_index.csv \
        --client "Northwind Systems Inc" --period-end 2025-12-31 \
        --out "Northwind - FY2025 PBC Status.xlsx"

--requests CSV:
    request_id, description, due_date, requested_from, keywords
    (keywords is a pipe-separated list of filename/path terms supplied by the
    client or engagement team, e.g. "bank statement|bank stmt|Nov 2025 bank")

--folder CSV:
    file_path, file_name
    (generate with e.g. `find . -type f -printf "%p,%f\n"` and add the header
    row, or any equivalent folder listing - this script does not scan disk)

--prior-year CSV (optional):
    prior_request_id, file_path
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")

REQUEST_COLS = ["request_id", "description", "due_date", "requested_from", "keywords"]
FOLDER_COLS = ["file_path", "file_name"]
PRIOR_COLS = ["prior_request_id", "file_path"]

STATUSES = {"matched", "gap", "no_keywords_supplied"}


def clean(s) -> str:
    return " ".join(str(s or "").split())


def load_csv(path: Path, cols: list[str]) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {path}")
    return [{k: clean(r.get(k)) for k in cols} for r in rows]


def write_template(path: Path, cols: list[str], sample: list[str]):
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        w.writerow(sample)
    print(f"Template written: {path}")


# ----------------------------------------------------------------------- match

def evaluate(requests: list[dict], folder: list[dict]):
    errors = []
    seen_ids = Counter(r["request_id"] for r in requests)
    dupes = [rid for rid, n in seen_ids.items() if n > 1]
    if dupes:
        errors.append(f"request_id appears more than once in the request list: "
                      f"{', '.join(sorted(set(dupes)))}. Each request must be unique.")

    results = []
    for i, req in enumerate(requests, start=1):
        rid = req["request_id"] or f"(blank at row {i})"
        if not req["request_id"]:
            errors.append(f"Row {i}: request_id is blank.")

        kw_raw = req["keywords"]
        if not kw_raw:
            results.append({**req, "status": "no_keywords_supplied", "matches": []})
            continue

        terms = [t.strip().lower() for t in kw_raw.split("|") if t.strip()]
        if not terms:
            results.append({**req, "status": "no_keywords_supplied", "matches": []})
            continue

        matches = []
        for doc in folder:
            name_l = doc["file_name"].lower()
            path_l = doc["file_path"].lower()
            for term in terms:
                if term in name_l or term in path_l:
                    matches.append({"file_path": doc["file_path"],
                                    "file_name": doc["file_name"],
                                    "keyword": term})
                    break  # one hit per document is enough to list it once

        status = "matched" if matches else "gap"
        results.append({**req, "status": status, "matches": matches})

    for r in results:
        if r["status"] not in STATUSES:
            errors.append(f"{r['request_id']}: resolved to unrecognized status "
                          f"'{r['status']}'.")

    result_ids = [r["request_id"] for r in results]
    input_ids = [r["request_id"] for r in requests]
    if Counter(result_ids) != Counter(input_ids):
        errors.append("Internal: output request list does not match input 1:1.")

    return results, errors


def load_prior(path: Path | None) -> dict:
    if not path or not path.exists():
        return {}
    rows = load_csv(path, PRIOR_COLS)
    idx: dict[str, list[str]] = {}
    for r in rows:
        idx.setdefault(r["prior_request_id"], []).append(r["file_path"])
    return idx


# ------------------------------------------------------------------------ xlsx

def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


def sheet_summary(wb, meta, results, errors, valid):
    ws = wb.active
    ws.title = "Workpaper Summary"
    widths(ws, {1: 4, 2: 46, 3: 60})
    r = 1

    def line(label, val=None, *, bold=False, size=11, fill=None):
        nonlocal r
        c = ws.cell(row=r, column=2, value=label)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        if val is not None:
            ws.cell(row=r, column=3, value=val)
        r += 1

    line("AUDIT PBC ASSEMBLER", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    v = ws.cell(row=r, column=2,
                value="ASSEMBLY STANDARD MET" if valid else
                "ASSEMBLY STANDARD NOT MET - see Exceptions tab")
    v.font = OK_FONT if valid else BAD_FONT
    if not valid:
        v.fill = BAD_FILL
    r += 2

    con = Counter(x["status"] for x in results)
    line("RESULTS", bold=True, size=12, fill=SUB_FILL)
    for k in ("matched", "gap", "no_keywords_supplied"):
        line(f"  {k}", con.get(k, 0))
    r += 1

    if errors:
        line("EXCEPTIONS", bold=True, size=12, fill=BAD_FILL)
        for e in errors:
            line("  " + e)
        r += 1

    line("NOTE", bold=True, size=11, fill=WARN_FILL)
    line("  A 'matched' status means candidate evidence was located by keyword. It does not "
         "mean the request is satisfied - the engagement team must confirm the document "
         "actually fulfills what was asked for.")
    r += 2
    line("Prepared by / date:  __________________  ____________", bold=True)
    line("Reviewed by / date:  __________________  ____________", bold=True)


def sheet_status(wb, results):
    ws = wb.create_sheet("PBC Status")
    heads = ["Request", "Description", "Status", "Matched file(s)", "Matching keyword(s)",
              "Due date", "Requested from"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    hdr(ws, len(heads))
    widths(ws, {1: 12, 2: 40, 3: 20, 4: 50, 5: 30, 6: 14, 7: 20})
    r = 2
    for x in results:
        ws.cell(row=r, column=1, value=x["request_id"])
        ws.cell(row=r, column=2, value=x["description"])
        sc = ws.cell(row=r, column=3, value=x["status"])
        if x["status"] == "matched":
            sc.font = OK_FONT
        elif x["status"] == "gap":
            sc.font = BAD_FONT
            for col in range(1, 8):
                ws.cell(row=r, column=col).fill = BAD_FILL
        else:
            sc.fill = WARN_FILL
        files = "; ".join(m["file_path"] for m in x["matches"])
        kws = "; ".join(sorted({m["keyword"] for m in x["matches"]}))
        ws.cell(row=r, column=4, value=files)
        ws.cell(row=r, column=5, value=kws)
        ws.cell(row=r, column=6, value=x["due_date"])
        ws.cell(row=r, column=7, value=x["requested_from"])
        r += 1
    ws.freeze_panes = "A2"


def sheet_gaps(wb, results):
    ws = wb.create_sheet("Gap List")
    heads = ["Request", "Description", "Status", "Due date", "Requested from"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    hdr(ws, len(heads))
    widths(ws, {1: 12, 2: 46, 3: 20, 4: 14, 5: 20})
    gaps = [x for x in results if x["status"] in ("gap", "no_keywords_supplied")]
    gaps.sort(key=lambda x: (x["due_date"] or "9999-99-99"))
    r = 2
    for x in gaps:
        ws.cell(row=r, column=1, value=x["request_id"])
        ws.cell(row=r, column=2, value=x["description"])
        c = ws.cell(row=r, column=3, value=x["status"])
        c.fill = BAD_FILL if x["status"] == "gap" else WARN_FILL
        ws.cell(row=r, column=4, value=x["due_date"])
        ws.cell(row=r, column=5, value=x["requested_from"])
        r += 1
    if not gaps:
        ws.cell(row=2, column=1, value="No gaps - every request has keywords and at least "
                "one candidate match.")
    ws.freeze_panes = "A2"


def sheet_prior_year(wb, prior_idx, results):
    ws = wb.create_sheet("Prior-Year Carryforward")
    heads = ["Request", "Description", "Current status", "Prior-year file(s)"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    hdr(ws, len(heads))
    widths(ws, {1: 12, 2: 46, 3: 20, 4: 60})
    r = 2
    by_id = {x["request_id"]: x for x in results}
    for rid, files in prior_idx.items():
        x = by_id.get(rid)
        ws.cell(row=r, column=1, value=rid)
        ws.cell(row=r, column=2, value=x["description"] if x else "(not in current request list)")
        ws.cell(row=r, column=3, value=x["status"] if x else "n/a")
        ws.cell(row=r, column=4, value="; ".join(files))
        r += 1
    if not prior_idx:
        ws.cell(row=2, column=1, value="No prior-year index supplied.")
    ws.freeze_panes = "A2"


def sheet_exceptions(wb, errors):
    ws = wb.create_sheet("Exceptions")
    ws.cell(row=1, column=1, value="Exception")
    hdr(ws, 1)
    widths(ws, {1: 120})
    for r, e in enumerate(errors, start=2):
        c = ws.cell(row=r, column=1, value=e)
        c.fill = BAD_FILL
        c.alignment = Alignment(wrap_text=True)
    if not errors:
        ws.cell(row=2, column=1, value="None.")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--requests")
    ap.add_argument("--folder")
    ap.add_argument("--prior-year")
    ap.add_argument("--out")
    ap.add_argument("--client", default="")
    ap.add_argument("--period-end", default="")
    ap.add_argument("--template-requests", help="emit a blank PBC request list template")
    ap.add_argument("--template-folder", help="emit a blank document folder listing template")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if args.template_requests:
        write_template(Path(args.template_requests), REQUEST_COLS,
                       ["REQ-001", "November 2025 bank statement", "2026-01-15", "Controller",
                        "bank statement|bank stmt|Nov 2025 bank"])
        return 0
    if args.template_folder:
        write_template(Path(args.template_folder), FOLDER_COLS,
                       ["/docs/banking/Nov2025_BankStmt.pdf", "Nov2025_BankStmt.pdf"])
        return 0

    if not args.requests or not args.folder or not args.out:
        sys.exit("--requests, --folder, and --out are required unless using a --template-* flag")

    requests = load_csv(Path(args.requests), REQUEST_COLS)
    folder = load_csv(Path(args.folder), FOLDER_COLS)
    prior_idx = load_prior(Path(args.prior_year) if args.prior_year else None)

    results, errors = evaluate(requests, folder)
    valid = not errors

    meta = {
        "Client": args.client or "(not stated)",
        "Period end": args.period_end or "(not stated)",
        "Requests": len(requests),
        "Documents in folder listing": len(folder),
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    print("=" * 76)
    print("AUDIT PBC ASSEMBLER")
    print("=" * 76)
    print(f"Client: {meta['Client']}    Requests: {len(requests)}    "
          f"Documents: {len(folder)}")
    con = Counter(x["status"] for x in results)
    for k in ("matched", "gap", "no_keywords_supplied"):
        print(f"  {k}: {con.get(k, 0)}")
    if errors:
        print(f"\nEXCEPTIONS ({len(errors)}):")
        for e in errors[:25]:
            print(f"  ! {e}")
        if len(errors) > 25:
            print(f"  ... and {len(errors) - 25} more")

    if not valid and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("The request list has a structural problem (duplicate or blank request_id). "
              "Fix the request list and re-run, or re-run with --force to document the gap.")
        print("=" * 76)
        return 1

    wb = Workbook()
    sheet_summary(wb, meta, results, errors, valid)
    sheet_status(wb, results)
    sheet_gaps(wb, results)
    sheet_prior_year(wb, prior_idx, results)
    sheet_exceptions(wb, errors)
    wb.active = 0

    out = Path(args.out)
    if not valid:
        out = out.with_name(out.stem + " [GATE NOT MET]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'ASSEMBLY STANDARD MET' if valid else 'NOT MET (forced)'}.  Workbook: {out}")
    return 0 if valid else 1


if __name__ == "__main__":
    sys.exit(main())

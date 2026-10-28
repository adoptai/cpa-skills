#!/usr/bin/env python3
"""
Recompute ASC 842 lease amortization schedules and tie ROU asset and lease
liability balances to the GL.

Gates:
  * incremental_borrowing_rate, term_months, payment_amount, and classification
    must be populated for every lease - never assumed or defaulted to a round
    number. A blank or unparseable field is a visible exception, not a plug.
  * any lease with a modification_date is NOT mechanically recomputed through
    the modification - remeasurement is a professional determination. It is
    routed whole to the Modification Flags tab.
  * for every non-modified lease, the recomputed ending lease liability and ROU
    asset must tie to the GL balance supplied for that lease, or the script
    reports a named, dollar-quantified difference. Nothing is plugged.

Usage - emit a blank lease register template:
    python3 lease_schedule.py --template --out lease_register.csv

Usage - run the tie-out:
    python3 lease_schedule.py --leases lease_register.csv --gl gl_balances.csv \
        --period-end 2025-12-31 --client "Northwind Systems Inc" \
        --out "Northwind - FY2025 Lease Tie-Out.xlsx"

--leases CSV:
    lease_id, description, commencement_date, term_months, payment_amount,
    payment_frequency (monthly|quarterly|annual), incremental_borrowing_rate
    (annual rate as a decimal, e.g. 0.06 - taken from the lease agreement or
    amortization table the client already uses, never assumed by this script),
    classification (operating|finance), modification_date (optional, blank if
    none)

--gl CSV:
    lease_id, rou_asset_balance, lease_liability_balance
    (period-end GL balances for the lease's ROU asset and lease liability)
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
CENT = Decimal("0.01")
MONEY = '#,##0.00;[Red](#,##0.00)'
PCT = '0.0"%"'

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")

CLASSES = {"operating", "finance"}
FREQ_PERIODS_PER_YEAR = {"monthly": 12, "quarterly": 4, "annual": 1}

LEASE_COLS = ["lease_id", "description", "commencement_date", "term_months",
              "payment_amount", "payment_frequency", "incremental_borrowing_rate",
              "classification", "modification_date"]
GL_COLS = ["lease_id", "rou_asset_balance", "lease_liability_balance"]


def dec(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").replace("%", "").strip()
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable number: {raw!r}")
    return -v if neg else v


def clean(s) -> str:
    return " ".join(str(s or "").split())


def pdate(raw):
    s = clean(raw)
    if not s:
        return None
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s.split(" ")[0], f).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date {raw!r}. Prefer ISO YYYY-MM-DD.")


def pint(raw):
    s = clean(raw)
    if not s:
        return None
    try:
        return int(Decimal(s))
    except InvalidOperation:
        raise ValueError(f"Unparseable integer: {raw!r}")


def load_csv(path: Path, cols: list[str]) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {path}")
    return [{k: clean(r.get(k)) for k in cols} for r in rows]


def write_template(path: Path):
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(LEASE_COLS)
        w.writerow(["L-001", "Office space - 3rd floor", "2024-01-01", "60", "10000",
                    "monthly", "0.06", "operating", ""])
    print(f"Template written: {path}")


# --------------------------------------------------------------------- compute

def months_between(d1: date, d2: date) -> int:
    """Whole calendar months from d1 to d2, counting a month only once d2's day
    has reached d1's day (so a lease that commenced on the 15th has not yet
    completed a period on the 10th of the following month)."""
    months = (d2.year - d1.year) * 12 + (d2.month - d1.month)
    if d2.day < d1.day:
        months -= 1
    return months


def build_schedule(lease: dict, period_end: date | None
                    ) -> tuple[list[dict], Decimal, Decimal, list[str]]:
    """Return (periods elapsed through period_end, ending_liability,
    ending_rou, errors) for one lease. The full schedule is computed through
    lease end to get the amortization pattern right, but only periods that
    have actually elapsed by period_end are returned and tied to the GL -
    tying the GL to the lease-end balance would be meaningless for a lease
    still in progress."""
    errs = []
    n = FREQ_PERIODS_PER_YEAR.get(lease["payment_frequency"])
    if n is None:
        errs.append(f"unrecognized payment_frequency '{lease['payment_frequency']}' "
                     f"(expected monthly, quarterly, or annual)")
        return [], None, None, errs

    term = lease["term_months_n"]
    payment = lease["payment_amount_n"]
    rate = lease["incremental_borrowing_rate_n"]
    months_per_period = 12 // n
    periods_total = round(term / (12 / n))
    if periods_total <= 0:
        errs.append("term_months resolves to zero or negative periods at this frequency")
        return [], None, None, errs

    commencement = lease.get("commencement_date_d")
    if period_end is None or commencement is None:
        errs.append("period_end and commencement_date are both required to determine how "
                     "many periods have elapsed")
        return [], None, None, errs

    elapsed_months = months_between(commencement, period_end)
    if elapsed_months < 0:
        errs.append(f"period_end {period_end} is before commencement {commencement} - "
                     f"the lease has not started")
        return [], None, None, errs

    periods_elapsed = min(elapsed_months // months_per_period, periods_total)
    if periods_elapsed <= 0:
        errs.append(f"no full period has elapsed between commencement {commencement} and "
                     f"period end {period_end} - nothing to tie out yet")
        return [], None, None, errs

    periodic_rate = (rate / Decimal(n))

    # PV of an ordinary annuity: liability at commencement = PV of all payments,
    # discounted at the periodic rate derived from the supplied annual rate.
    pv = ZERO
    for p in range(1, periods_total + 1):
        pv += payment / ((1 + periodic_rate) ** p)
    pv = pv.quantize(CENT, rounding=ROUND_HALF_UP)

    total_cost = (payment * periods_total).quantize(CENT, rounding=ROUND_HALF_UP)
    straight_line_rou_amort = (pv / periods_total).quantize(CENT, rounding=ROUND_HALF_UP) \
        if lease["classification"] == "finance" else \
        (total_cost / periods_total).quantize(CENT, rounding=ROUND_HALF_UP)

    rows = []
    bal = pv
    rou = pv
    accrued_rounding = ZERO
    for p in range(1, periods_total + 1):
        interest = (bal * periodic_rate).quantize(CENT, rounding=ROUND_HALF_UP)
        principal = payment - interest
        if p == periods_total:
            # last period absorbs rounding so the schedule foots to exactly zero
            principal = bal
            interest = payment - principal
        end_bal = (bal - principal).quantize(CENT, rounding=ROUND_HALF_UP)

        if lease["classification"] == "operating":
            # single straight-line lease cost: ROU amortization is the plug that
            # makes straight-line expense (interest + amortization) equal each
            # period, floored at zero and never negative.
            amort = (total_cost / periods_total - interest).quantize(CENT, rounding=ROUND_HALF_UP)
            amort = max(amort, ZERO)
        else:
            amort = straight_line_rou_amort
        end_rou = (rou - amort).quantize(CENT, rounding=ROUND_HALF_UP)
        if p == periods_total:
            end_rou = ZERO

        rows.append({"period": p, "beg_liability": bal, "interest": interest,
                     "payment": payment, "principal": principal, "end_liability": end_bal,
                     "beg_rou": rou, "amortization": amort, "end_rou": end_rou})
        bal, rou = end_bal, end_rou

    elapsed_rows = rows[:periods_elapsed]
    return (elapsed_rows, elapsed_rows[-1]["end_liability"],
            elapsed_rows[-1]["end_rou"], errs)


# ------------------------------------------------------------------------ main

def evaluate(leases: list[dict], gl_by_id: dict, as_of: date):
    errors, modifications, results = [], [], []

    for lease in leases:
        lid = lease["lease_id"] or "(blank lease_id)"
        missing = []
        lease["incremental_borrowing_rate_n"] = None
        lease["term_months_n"] = None
        lease["payment_amount_n"] = None
        lease["commencement_date_d"] = None

        try:
            lease["commencement_date_d"] = pdate(lease["commencement_date"])
        except ValueError as e:
            errors.append(f"{lid}: commencement_date - {e}")
        if lease["commencement_date_d"] is None:
            missing.append("commencement_date")

        try:
            lease["incremental_borrowing_rate_n"] = dec(lease["incremental_borrowing_rate"])
        except ValueError as e:
            errors.append(f"{lid}: incremental_borrowing_rate - {e}")
        if lease["incremental_borrowing_rate_n"] is None:
            missing.append("incremental_borrowing_rate")

        try:
            lease["term_months_n"] = pint(lease["term_months"])
        except ValueError as e:
            errors.append(f"{lid}: term_months - {e}")
        if not lease["term_months_n"]:
            missing.append("term_months")

        try:
            lease["payment_amount_n"] = dec(lease["payment_amount"])
        except ValueError as e:
            errors.append(f"{lid}: payment_amount - {e}")
        if lease["payment_amount_n"] is None:
            missing.append("payment_amount")

        if lease["classification"] not in CLASSES:
            missing.append("classification")

        if missing:
            errors.append(f"{lid}: missing or unparseable required field(s): "
                          f"{', '.join(missing)}. A lease is not tied out on an assumed term, "
                          f"rate, payment, or classification.")
            results.append({**lease, "status": "exception", "periods": [],
                            "end_liability": None, "end_rou": None, "gl_liability": None,
                            "gl_rou": None, "diff_liability": None, "diff_rou": None})
            continue

        if lease["modification_date"]:
            modifications.append(lease)
            results.append({**lease, "status": "modification - not recomputed",
                            "periods": [], "end_liability": None, "end_rou": None,
                            "gl_liability": None, "gl_rou": None,
                            "diff_liability": None, "diff_rou": None})
            continue

        periods, end_liability, end_rou, errs = build_schedule(lease, as_of)
        if errs:
            for e in errs:
                errors.append(f"{lid}: {e}")
            results.append({**lease, "status": "exception", "periods": [],
                            "end_liability": None, "end_rou": None, "gl_liability": None,
                            "gl_rou": None, "diff_liability": None, "diff_rou": None})
            continue

        gl = gl_by_id.get(lid)
        gl_liability = dec(gl["lease_liability_balance"]) if gl else None
        gl_rou = dec(gl["rou_asset_balance"]) if gl else None
        diff_liability = (end_liability - gl_liability) if gl_liability is not None else None
        diff_rou = (end_rou - gl_rou) if gl_rou is not None else None

        if gl is None:
            errors.append(f"{lid}: no GL balance supplied - cannot tie out.")
            status = "no GL balance"
        elif diff_liability != ZERO or diff_rou != ZERO:
            errors.append(f"{lid}: does not tie. Liability recomputed {end_liability:,.2f} "
                          f"vs GL {gl_liability:,.2f} (diff {diff_liability:,.2f}); ROU "
                          f"recomputed {end_rou:,.2f} vs GL {gl_rou:,.2f} (diff "
                          f"{diff_rou:,.2f}).")
            status = "exception - does not tie"
        else:
            status = "tied"

        results.append({**lease, "status": status, "periods": periods,
                        "end_liability": end_liability, "end_rou": end_rou,
                        "gl_liability": gl_liability, "gl_rou": gl_rou,
                        "diff_liability": diff_liability, "diff_rou": diff_rou})

    return results, modifications, errors


def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


def money_cell(ws, r, c, v):
    cell = ws.cell(row=r, column=c, value=float(v) if v is not None else None)
    if v is not None:
        cell.number_format = MONEY
    return cell


def sheet_summary(wb, meta, results, errors, valid):
    ws = wb.active
    ws.title = "Workpaper Summary"
    widths(ws, {1: 4, 2: 46, 3: 20, 4: 20, 5: 64})
    r = 1

    def line(label, a=None, b=None, *, bold=False, size=11, fill=None):
        nonlocal r
        c = ws.cell(row=r, column=2, value=label)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        for col, v in ((3, a), (4, b)):
            if v is None:
                continue
            ws.cell(row=r, column=col, value=v)
        r += 1

    line("LEASE SCHEDULE TIE-OUT (ASC 842)", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    v = ws.cell(row=r, column=2,
                value="TIE-OUT STANDARD MET" if valid else
                "TIE-OUT STANDARD NOT MET - see Exceptions tab")
    v.font = OK_FONT if valid else BAD_FONT
    if not valid:
        v.fill = BAD_FILL
    r += 2

    con = Counter(x["status"] for x in results)
    line("RESULTS", bold=True, size=12, fill=SUB_FILL)
    for k, n in con.items():
        line(f"  {k}", n)
    r += 1

    if errors:
        line("EXCEPTIONS", bold=True, size=12, fill=BAD_FILL)
        for e in errors:
            line("  " + e)
        r += 1

    r += 1
    line("CONCLUSION", bold=True, size=12)
    line("  [Complete after review. State whether the recomputed schedules support the "
         "recorded ROU asset and lease liability balances, and whether any modification "
         "requires a remeasurement determination.]")
    r += 2
    line("Prepared by / date:  __________________  ____________", bold=True)
    line("Reviewed by / date:  __________________  ____________", bold=True)


def sheet_schedule(wb, results):
    ws = wb.create_sheet("Schedule")
    heads = ["Lease", "Period", "Beg liability", "Interest", "Payment", "Principal",
              "End liability", "Beg ROU", "Amortization", "End ROU"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    hdr(ws, len(heads))
    widths(ws, {1: 12, 2: 8, 3: 16, 4: 14, 5: 14, 6: 14, 7: 16, 8: 16, 9: 14, 10: 16})
    r = 2
    for x in results:
        for p in x["periods"]:
            ws.cell(row=r, column=1, value=x["lease_id"])
            ws.cell(row=r, column=2, value=p["period"])
            money_cell(ws, r, 3, p["beg_liability"])
            money_cell(ws, r, 4, p["interest"])
            money_cell(ws, r, 5, p["payment"])
            money_cell(ws, r, 6, p["principal"])
            money_cell(ws, r, 7, p["end_liability"])
            money_cell(ws, r, 8, p["beg_rou"])
            money_cell(ws, r, 9, p["amortization"])
            money_cell(ws, r, 10, p["end_rou"])
            r += 1
    ws.freeze_panes = "A2"


def sheet_tieout(wb, results):
    ws = wb.create_sheet("GL Tie-Out")
    heads = ["Lease", "Description", "Classification", "Status", "Recomputed liability",
              "GL liability", "Diff liability", "Recomputed ROU", "GL ROU", "Diff ROU"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    hdr(ws, len(heads))
    widths(ws, {1: 12, 2: 28, 3: 14, 4: 24, 5: 18, 6: 16, 7: 14, 8: 16, 9: 16, 10: 14})
    r = 2
    for x in results:
        if x["status"] in ("modification - not recomputed", "exception"):
            continue
        ws.cell(row=r, column=1, value=x["lease_id"])
        ws.cell(row=r, column=2, value=x["description"])
        ws.cell(row=r, column=3, value=x["classification"])
        sc = ws.cell(row=r, column=4, value=x["status"])
        tied = x["status"] == "tied"
        sc.font = OK_FONT if tied else BAD_FONT
        if not tied:
            for col in range(1, 11):
                ws.cell(row=r, column=col).fill = BAD_FILL
        money_cell(ws, r, 5, x["end_liability"])
        money_cell(ws, r, 6, x["gl_liability"])
        money_cell(ws, r, 7, x["diff_liability"])
        money_cell(ws, r, 8, x["end_rou"])
        money_cell(ws, r, 9, x["gl_rou"])
        money_cell(ws, r, 10, x["diff_rou"])
        r += 1
    ws.freeze_panes = "A2"


def sheet_modifications(wb, modifications):
    ws = wb.create_sheet("Modification Flags")
    heads = ["Lease", "Description", "Modification date", "Term (months)", "Payment",
              "Frequency", "IBR", "Classification", "Note"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    hdr(ws, len(heads))
    widths(ws, {1: 12, 2: 28, 3: 16, 4: 12, 5: 14, 6: 12, 7: 10, 8: 14, 9: 60})
    r = 2
    for x in modifications:
        ws.cell(row=r, column=1, value=x["lease_id"])
        ws.cell(row=r, column=2, value=x["description"])
        ws.cell(row=r, column=3, value=x["modification_date"])
        ws.cell(row=r, column=4, value=x["term_months"])
        ws.cell(row=r, column=5, value=x["payment_amount"])
        ws.cell(row=r, column=6, value=x["payment_frequency"])
        ws.cell(row=r, column=7, value=x["incremental_borrowing_rate"])
        ws.cell(row=r, column=8, value=x["classification"])
        n = ws.cell(row=r, column=9, value="Not mechanically recomputed - remeasurement "
                    "from the modification date requires a professional determination of the "
                    "revised discount rate, term, and payments. Route to the preparer.")
        n.fill = WARN_FILL
        r += 1
    if not modifications:
        ws.cell(row=2, column=1, value="No leases flagged with a modification date.")
    ws.freeze_panes = "A2"


def sheet_exceptions(wb, errors):
    ws = wb.create_sheet("Exceptions")
    ws.cell(row=1, column=1, value="Exception")
    hdr(ws, 1)
    widths(ws, {1: 120})
    for r, e in enumerate(errors, start=2):
        c = ws.cell(row=r, column=1, value=e)
        c.fill = BAD_FILL
        c.alignment = Alignment(wrap_text=True)
    if not errors:
        ws.cell(row=2, column=1, value="None.")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--leases")
    ap.add_argument("--gl")
    ap.add_argument("--out", required=True)
    ap.add_argument("--period-end")
    ap.add_argument("--client", default="")
    ap.add_argument("--template", action="store_true",
                    help="emit a blank lease register instead of running the tie-out")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if args.template:
        write_template(Path(args.out))
        return 0

    if not args.leases or not args.gl or not args.period_end:
        sys.exit("--leases, --gl, and --period-end are required unless --template is given")

    leases = load_csv(Path(args.leases), LEASE_COLS)
    gl_rows = load_csv(Path(args.gl), GL_COLS)
    gl_by_id = {g["lease_id"]: g for g in gl_rows}
    period_end = pdate(args.period_end)

    results, modifications, errors = evaluate(leases, gl_by_id, period_end)
    valid = not errors

    meta = {
        "Client": args.client or "(not stated)",
        "Period end": str(period_end) if period_end else "(not stated)",
        "Leases in register": len(leases),
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    print("=" * 76)
    print("LEASE SCHEDULE TIE-OUT (ASC 842)")
    print("=" * 76)
    print(f"Client: {meta['Client']}    Period end: {meta['Period end']}")
    con = Counter(x["status"] for x in results)
    for k, n in con.items():
        print(f"  {k}: {n}")
    if modifications:
        print(f"\nMODIFICATIONS FLAGGED, NOT RECOMPUTED ({len(modifications)}):")
        for m in modifications:
            print(f"  ! {m['lease_id']}: modified {m['modification_date']} - "
                  f"route for remeasurement")
    if errors:
        print(f"\nEXCEPTIONS ({len(errors)}):")
        for e in errors[:25]:
            print(f"  ! {e}")
        if len(errors) > 25:
            print(f"  ... and {len(errors) - 25} more")

    if not valid and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("One or more leases are missing required terms or do not tie to the GL. "
              "Resolve the exceptions above, or re-run with --force to document the gap.")
        print("=" * 76)
        return 1

    wb = Workbook()
    sheet_summary(wb, meta, results, errors, valid)
    sheet_schedule(wb, results)
    sheet_tieout(wb, results)
    sheet_modifications(wb, modifications)
    sheet_exceptions(wb, errors)
    wb.active = 0

    out = Path(args.out)
    if not valid:
        out = out.with_name(out.stem + " [GATE NOT MET]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'TIE-OUT STANDARD MET' if valid else 'NOT MET (forced)'}.  Workbook: {out}")
    return 0 if valid else 1


if __name__ == "__main__":
    sys.exit(main())

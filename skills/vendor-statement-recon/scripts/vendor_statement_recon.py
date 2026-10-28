#!/usr/bin/env python3
"""
Reconcile vendor statements to the AP subledger, per vendor: classify every
statement line and every open subledger line, and tie the reconciled total to
the subledger open balance - or produce a quantified, named difference.

Gates:
  * every statement line AND every open subledger line for a vendor lands in
    exactly one classification bucket - nothing dropped silently
  * an unparseable amount or date is a visible, blocking exception - never
    coerced or dropped
  * a duplicate (vendor, invoice_no) key in the subledger makes matching
    unreliable and blocks the workbook until resolved
  * the reconciling total is never forced to zero - a real difference is
    reported as a named, quantified unreconciled amount
  * "proposed action" per discrepancy is clearly labeled a suggestion only;
    the skill never concludes a discrepancy is resolved

Usage - emit blank templates:
    python3 vendor_statement_recon.py --template --out statement_template.csv

Usage - reconcile:
    python3 vendor_statement_recon.py --statement vendor_statements.csv \
        --subledger ap_subledger.csv --client "Northwind Systems Inc" \
        --out "Northwind - Vendor Statement Recon.xlsx"

--statement CSV (vendor statement extract, all vendors together):
    vendor, statement_date, ref, amount, line_type
    (line_type: invoice | credit | payment)

--subledger CSV (AP subledger export):
    vendor, invoice_no, amount, invoice_date, status
    (status: open | paid)
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  pip install openpyxl")

ZERO = Decimal("0.00")
TOLERANCE = Decimal("0.01")
MONEY = '#,##0.00;[Red](#,##0.00)'
DATEF = "yyyy-mm-dd"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")

STMT_COLS = ["vendor", "statement_date", "ref", "amount", "line_type"]
SUB_COLS = ["vendor", "invoice_no", "amount", "invoice_date", "status"]
LINE_TYPES = {"invoice", "credit", "payment"}

SUGGESTED_ACTION = {
    "missing-from-subledger-invoice": "SUGGESTED (not a determination): request invoice copy and post to subledger",
    "missing-from-subledger-credit": "SUGGESTED (not a determination): request credit memo copy and apply",
    "missing-from-subledger-payment": "SUGGESTED (not a determination): confirm payment applied; trace to bank",
    "disputed": "SUGGESTED (not a determination): compare source documents to resolve the amount difference",
    "in-subledger-not-on-statement": "SUGGESTED (not a determination): confirm with vendor whether item is stale or statement is incomplete",
}


def dec(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def clean(s) -> str:
    return " ".join(str(s or "").split())


def pdate(raw):
    s = clean(raw)
    if not s:
        return None
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s.split(" ")[0], f).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date {raw!r}. Prefer ISO YYYY-MM-DD.")


def write_template(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=STMT_COLS)
        w.writeheader()
        w.writerow({"vendor": "Acme Supply Co", "statement_date": "2026-01-31",
                    "ref": "9981", "amount": "1250.00", "line_type": "invoice"})
    sub_path = path.with_name(path.stem + "_subledger" + path.suffix)
    with sub_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=SUB_COLS)
        w.writeheader()
        w.writerow({"vendor": "Acme Supply Co", "invoice_no": "9981", "amount": "1250.00",
                    "invoice_date": "2026-01-15", "status": "open"})
    print(f"Templates written: {path} and {sub_path}")


def load_statement(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        missing = [c for c in STMT_COLS if c not in (reader.fieldnames or [])]
        if missing:
            sys.exit(f"{path}: missing required column(s) {missing}")
        raw_rows = [{k: clean(r.get(k)) for k in STMT_COLS} for r in reader]
    if not raw_rows:
        sys.exit(f"No rows in {path}")
    rows, errors = [], []
    for i, r in enumerate(raw_rows, start=2):
        r["row"] = i
        lt = r["line_type"].lower()
        if lt not in LINE_TYPES:
            errors.append(f"statement row {i}: unrecognized line_type {r['line_type']!r} (expected invoice/credit/payment)")
            continue
        r["line_type"] = lt
        try:
            amt = dec(r["amount"])
            if amt is None:
                raise ValueError("blank amount")
            r["amount_dec"] = abs(amt)
        except ValueError as e:
            errors.append(f"statement row {i}: {e}")
            continue
        try:
            r["statement_date_dec"] = pdate(r["statement_date"])
        except ValueError as e:
            errors.append(f"statement row {i}: {e}")
            continue
        rows.append(r)
    return rows, errors


def load_subledger(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        missing = [c for c in SUB_COLS if c not in (reader.fieldnames or [])]
        if missing:
            sys.exit(f"{path}: missing required column(s) {missing}")
        raw_rows = [{k: clean(r.get(k)) for k in SUB_COLS} for r in reader]
    if not raw_rows:
        sys.exit(f"No rows in {path}")
    rows, errors = [], []
    seen_keys = {}
    for i, r in enumerate(raw_rows, start=2):
        r["row"] = i
        status = r["status"].lower()
        if status not in ("open", "paid"):
            errors.append(f"subledger row {i}: unrecognized status {r['status']!r} (expected open/paid)")
            continue
        r["status"] = status
        try:
            amt = dec(r["amount"])
            if amt is None:
                raise ValueError("blank amount")
            r["amount_dec"] = abs(amt)
        except ValueError as e:
            errors.append(f"subledger row {i}: {e}")
            continue
        try:
            r["invoice_date_dec"] = pdate(r["invoice_date"])
        except ValueError as e:
            errors.append(f"subledger row {i}: {e}")
            continue
        key = (r["vendor"], r["invoice_no"])
        if key in seen_keys:
            errors.append(
                f"subledger row {i}: (vendor, invoice_no) {key} duplicates row "
                f"{seen_keys[key]} - match target is ambiguous"
            )
            continue
        seen_keys[key] = i
        rows.append(r)
    return rows, errors


def reconcile(statement_rows, subledger_rows):
    by_vendor_stmt = defaultdict(list)
    for r in statement_rows:
        by_vendor_stmt[r["vendor"]].append(r)
    by_vendor_sub = defaultdict(list)
    for r in subledger_rows:
        by_vendor_sub[r["vendor"]].append(r)

    vendors = sorted(set(by_vendor_stmt) | set(by_vendor_sub))
    vendor_summaries, discrepancies = [], []

    for vendor in vendors:
        stmt_lines = by_vendor_stmt.get(vendor, [])
        sub_lines = by_vendor_sub.get(vendor, [])
        sub_by_invno = {r["invoice_no"]: r for r in sub_lines}
        referenced_invnos = set()

        matched_n = disputed_n = missing_n = 0
        statement_balance = ZERO
        adjustment = ZERO

        for line in stmt_lines:
            sign = Decimal(1) if line["line_type"] in ("invoice",) else Decimal(-1)
            statement_balance += sign * line["amount_dec"]

            sub = sub_by_invno.get(line["ref"])
            if sub is None:
                missing_n += 1
                bucket = f"missing-from-subledger-{line['line_type']}"
                discrepancies.append({
                    "vendor": vendor, "type": "missing-from-subledger",
                    "ref": line["ref"], "line_type": line["line_type"],
                    "amount": line["amount_dec"], "detail": "on statement, not found in subledger",
                    "action": SUGGESTED_ACTION[bucket],
                })
                adjustment += sign * line["amount_dec"]
                continue
            referenced_invnos.add(line["ref"])
            if abs(sub["amount_dec"] - line["amount_dec"]) <= TOLERANCE:
                matched_n += 1
            else:
                disputed_n += 1
                discrepancies.append({
                    "vendor": vendor, "type": "disputed",
                    "ref": line["ref"], "line_type": line["line_type"],
                    "amount": line["amount_dec"],
                    "detail": f"statement {line['amount_dec']} vs subledger {sub['amount_dec']}",
                    "action": SUGGESTED_ACTION["disputed"],
                })

        not_on_statement = [
            r for r in sub_lines if r["status"] == "open" and r["invoice_no"] not in referenced_invnos
        ]
        for r in not_on_statement:
            discrepancies.append({
                "vendor": vendor, "type": "in-subledger-not-on-statement",
                "ref": r["invoice_no"], "line_type": "invoice",
                "amount": r["amount_dec"], "detail": "open in subledger, not referenced on statement",
                "action": SUGGESTED_ACTION["in-subledger-not-on-statement"],
            })
            adjustment -= r["amount_dec"]

        subledger_open_balance = sum((r["amount_dec"] for r in sub_lines if r["status"] == "open"), ZERO)
        recon_balance = subledger_open_balance + adjustment
        difference = statement_balance - recon_balance
        tie = abs(difference) <= TOLERANCE

        vendor_summaries.append({
            "vendor": vendor,
            "statement_lines": len(stmt_lines),
            "subledger_open_lines": sum(1 for r in sub_lines if r["status"] == "open"),
            "matched": matched_n, "disputed": disputed_n, "missing": missing_n,
            "not_on_statement": len(not_on_statement),
            "statement_balance": statement_balance,
            "subledger_open_balance": subledger_open_balance,
            "recon_balance": recon_balance,
            "difference": difference,
            "tie": tie,
        })

    return vendor_summaries, discrepancies


def style_header(ws, cols):
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def autosize(ws, cols):
    for c, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(len(name) + 2, 14)


def sheet_per_vendor(wb, summaries):
    ws = wb.active
    ws.title = "Per-Vendor Reconciliation"
    cols = ["vendor", "statement_lines", "subledger_open_lines", "matched", "disputed",
            "missing", "not_on_statement", "statement_balance", "subledger_open_balance",
            "recon_balance", "difference", "tie"]
    style_header(ws, cols)
    for r, s in enumerate(summaries, start=2):
        ws.cell(row=r, column=1, value=s["vendor"])
        ws.cell(row=r, column=2, value=s["statement_lines"])
        ws.cell(row=r, column=3, value=s["subledger_open_lines"])
        ws.cell(row=r, column=4, value=s["matched"])
        ws.cell(row=r, column=5, value=s["disputed"])
        ws.cell(row=r, column=6, value=s["missing"])
        ws.cell(row=r, column=7, value=s["not_on_statement"])
        for col, key in ((8, "statement_balance"), (9, "subledger_open_balance"),
                         (10, "recon_balance"), (11, "difference")):
            c = ws.cell(row=r, column=col, value=float(s[key]))
            c.number_format = MONEY
        tc = ws.cell(row=r, column=12, value="TIE" if s["tie"] else "DIFFERENCE")
        tc.font = OK_FONT if s["tie"] else BAD_FONT
        if not s["tie"]:
            for c in range(1, 13):
                ws.cell(row=r, column=c).fill = BAD_FILL
    autosize(ws, cols)


def sheet_discrepancies(wb, discrepancies):
    ws = wb.create_sheet("Discrepancy List")
    cols = ["vendor", "type", "ref", "line_type", "amount", "detail", "action"]
    style_header(ws, cols)
    for r, d in enumerate(discrepancies, start=2):
        ws.cell(row=r, column=1, value=d["vendor"])
        ws.cell(row=r, column=2, value=d["type"])
        ws.cell(row=r, column=3, value=d["ref"])
        ws.cell(row=r, column=4, value=d["line_type"])
        ac = ws.cell(row=r, column=5, value=float(d["amount"])); ac.number_format = MONEY
        ws.cell(row=r, column=6, value=d["detail"])
        ws.cell(row=r, column=7, value=d["action"])
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = WARN_FILL
    autosize(ws, cols)


def sheet_unreconciled(wb, summaries, hard_errors):
    ws = wb.create_sheet("Unreconciled-Exceptions")
    cols = ["vendor", "statement_balance", "recon_balance", "difference"]
    style_header(ws, cols)
    r = 2
    for s in summaries:
        if s["tie"]:
            continue
        ws.cell(row=r, column=1, value=s["vendor"])
        for col, key in ((2, "statement_balance"), (3, "recon_balance"), (4, "difference")):
            c = ws.cell(row=r, column=col, value=float(s[key]))
            c.number_format = MONEY
        for c in range(1, 5):
            ws.cell(row=r, column=c).fill = BAD_FILL
        r += 1
    if hard_errors:
        r += 1
        ws.cell(row=r, column=1, value="-- gate failures (blocking) --").font = BAD_FONT
        r += 1
        for e in hard_errors:
            cell = ws.cell(row=r, column=1, value=e)
            cell.fill = BAD_FILL
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            r += 1
    autosize(ws, cols)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--statement")
    ap.add_argument("--subledger")
    ap.add_argument("--out", required=True)
    ap.add_argument("--client", default="")
    ap.add_argument("--template", action="store_true")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if args.template:
        write_template(Path(args.out))
        return 0

    if not (args.statement and args.subledger):
        sys.exit("--statement and --subledger are both required (unless --template)")

    stmt_rows, stmt_errors = load_statement(Path(args.statement))
    sub_rows, sub_errors = load_subledger(Path(args.subledger))
    errors = stmt_errors + sub_errors
    valid = not errors

    summaries, discrepancies = reconcile(stmt_rows, sub_rows) if (valid or args.force) else ([], [])
    n_tied = sum(1 for s in summaries if s["tie"])
    n_diff = len(summaries) - n_tied

    meta = {
        "Client": args.client or "(not stated)",
        "Vendors": len(summaries),
        "Statement lines read": len(stmt_rows) + len(stmt_errors),
        "Subledger lines read": len(sub_rows) + len(sub_errors),
        "Prepared": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    print("=" * 76)
    print("VENDOR STATEMENT RECONCILIATION")
    print("=" * 76)
    print(f"Client            : {meta['Client']}")
    print(f"Vendors           : {len(summaries)}   tied {n_tied}   difference {n_diff}")
    print(f"Discrepancies     : {len(discrepancies)}")
    if errors:
        print(f"\nGATE FAILURES ({len(errors)}):")
        for e in errors[:25]:
            print(f"  ! {e}")
        if len(errors) > 25:
            print(f"  ... and {len(errors) - 25} more")

    if not valid and not args.force:
        print("\n" + "=" * 76)
        print("WORKBOOK NOT WRITTEN.")
        print("Every amount and date must parse, and (vendor, invoice_no) must be unique")
        print("in the subledger before matching is reliable. Fix the flagged rows and")
        print("re-run.")
        print("=" * 76)
        return 1

    wb = Workbook()
    sheet_per_vendor(wb, summaries)
    sheet_discrepancies(wb, discrepancies)
    sheet_unreconciled(wb, summaries, errors if args.force else [])
    wb.active = 0
    out = Path(args.out)
    if not valid:
        out = out.with_name(out.stem + " [GATE NOT MET]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n{'GATE CLEAN' if valid else 'NOT MET (forced)'}.  Workbook: {out}")
    return 0 if valid else 1


if __name__ == "__main__":
    sys.exit(main())
